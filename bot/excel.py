"""openpyxl workbook builders. All builders are synchronous and return the ``.xlsx`` bytes."""

from __future__ import annotations

import re
from collections.abc import Iterable, Sequence
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import FullProfile, Group, Student, Tutor
from .texts import course_label, residence_label, social_labels

StudentRow = Student
"""A student row: :class:`Student` already carries ``tutor_name`` and ``group_name``."""

ALL_SHEET_NAME = "Barchasi"
MAX_SHEET_NAME_LEN = 31
RESERVED_SHEET_NAME = "History"  # Excel keeps this name for its own change-history sheet

HEADERS: tuple[str, ...] = (
    "№",
    "F.I.SH",
    "Telefon",
    "Yo'nalish",
    "Tyutor",
    "Guruh",
    "Turar joy",
    "Manzil",
    "Otasining F.I.SH",
    "Otasining tel",
    "Onasining F.I.SH",
    "Onasining tel",
    "Telegram username",
    "Telegram ID",
    "Ro'yxatdan o'tgan vaqt",
    "Oxirgi tahrir",
)
COLUMN_WIDTHS: tuple[int, ...] = (5, 32, 16, 24, 26, 16, 12, 36, 32, 16, 32, 16, 20, 14, 20, 20)
_TEXT_COLUMNS = {3, 10, 12}  # 1-based indexes of phone columns; forced to text format

_INVALID_SHEET_CHARS_RE = re.compile(r"[\[\]:*?/\\]")
_SHEET_EDGE_RE = re.compile(r"^[\s']+|[\s']+$")  # Excel forbids a leading/trailing apostrophe
_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9E1F2")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def sanitize_sheet_name(name: str) -> str:
    """Make ``name`` a valid Excel sheet title: no ``[]:*?/\\``, no edge apostrophes, at most 31 chars."""
    cleaned = _SHEET_EDGE_RE.sub("", _INVALID_SHEET_CHARS_RE.sub("", name))
    # Truncation can expose a new trailing apostrophe (e.g. "... o'g'li" cut at "o'"), so strip again.
    cleaned = _SHEET_EDGE_RE.sub("", cleaned[:MAX_SHEET_NAME_LEN])
    if cleaned.lower() == RESERVED_SHEET_NAME.lower():
        cleaned += "_"
    return cleaned or "Sheet"


def unique_sheet_names(names: Iterable[str]) -> list[str]:
    """Sanitise names and make them unique (case-insensitively) by appending `` (2)``, `` (3)``, ..."""
    result: list[str] = []
    seen: set[str] = set()
    for raw in names:
        base = sanitize_sheet_name(raw)
        candidate = base
        counter = 2
        while candidate.lower() in seen:
            suffix = f" ({counter})"
            candidate = _SHEET_EDGE_RE.sub("", base[: MAX_SHEET_NAME_LEN - len(suffix)]) + suffix
            counter += 1
        seen.add(candidate.lower())
        result.append(candidate)
    return result


def _fill_sheet(ws: Worksheet, students: Sequence[StudentRow]) -> None:
    ws.append(list(HEADERS))
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _BORDER

    for index, s in enumerate(students, start=1):
        ws.append(
            [
                index,
                s.full_name,
                str(s.phone),
                s.direction,
                s.tutor_name,
                s.group_name,
                residence_label(s.residence),
                s.address,
                s.father_name,
                str(s.father_phone),
                s.mother_name,
                str(s.mother_phone),
                f"@{s.username}" if s.username else None,
                s.telegram_id,
                s.created_at,
                # NULL until the data is changed for the first time, so a filled cell means
                # "edited since registration" at a glance.
                s.edited_at,
            ]
        )
        row = ws[index + 1]
        for col_index, cell in enumerate(row, start=1):
            cell.border = _BORDER
            if col_index in _TEXT_COLUMNS:
                cell.number_format = "@"

    for col_index, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(ws.max_row, 1)}"


def build_students_workbook(sheets: Sequence[tuple[str, Sequence[StudentRow]]]) -> bytes:
    """Build a workbook with one sheet per ``(name, students)`` pair and return its bytes."""
    wb = Workbook()
    default = wb.active
    if default is not None:
        wb.remove(default)
    if not sheets:
        sheets = [(ALL_SHEET_NAME, [])]
    names = unique_sheet_names(name for name, _ in sheets)
    for title, (_, students) in zip(names, sheets, strict=True):
        ws = wb.create_sheet(title=title)
        _fill_sheet(ws, students)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_group_workbook(group: Group, students: Sequence[StudentRow]) -> bytes:
    """Single sheet named after the group."""
    return build_students_workbook([(group.name, students)])


def build_residence_workbook(residence: str, students: Sequence[StudentRow]) -> bytes:
    """Single sheet named after the residence label (rows keep their ``Guruh`` column)."""
    return build_students_workbook([(residence_label(residence), students)])


def build_tutor_workbook(groups: Sequence[Group], students: Sequence[StudentRow]) -> bytes:
    """Sheet ``Barchasi`` with all of the tutor's students + one sheet per group."""
    sheets: list[tuple[str, Sequence[StudentRow]]] = [(ALL_SHEET_NAME, students)]
    for group in groups:
        sheets.append((group.name, [s for s in students if s.group_id == group.id]))
    return build_students_workbook(sheets)


def build_all_tutors_workbook(tutors: Sequence[Tutor], students: Sequence[StudentRow]) -> bytes:
    """Sheet ``Barchasi`` with every student + one sheet per tutor."""
    sheets: list[tuple[str, Sequence[StudentRow]]] = [(ALL_SHEET_NAME, students)]
    for tutor in tutors:
        sheets.append((tutor.name, [s for s in students if s.tutor_id == tutor.id]))
    return build_students_workbook(sheets)


# ------------------------------------------------------- full survey workbook

FULL_HEADERS: tuple[str, ...] = (
    "№",
    "Tyutori F.I.Sh.",
    "Tyutor tel. raqami",
    "Talaba F.I.Sh.",
    "Talaba tel raqami",
    "Talim yo'nalishi",
    "Guruhi",
    "Kursi",
    "Pasport seriya raqami",
    "Pasport JShShR (PNFL)",
    "Tug'ilgan kun oy yili",
    "Fuqoroligi",
    "Viloyati",
    "Shahar (Tuman)",
    "MFY",
    "MFY raqami (MFY raisi, yoshlar yetakchisi)",
    "Ko'cha uy raqami",
    "Ish bilan bandligi (Ishlaydi, ishlamaydi)",
    "Ishlaydigan tashkiloti nomi",
    "Lavozimi",
    "Ishlaydigan tashkiloti joylashgan joyi",
    "Ishlaydigan tashkilot tel raqami",
    "Olila qurgan (Oila qurmagan)",
    "Turmush o'rtog'ini F.I.Sh.",
    "Turmush o'rtog'ining ish joyi",
    "Turmush o'rtog'ini tel raqami",
    "Ijtimoiy holati",
    "Otasini F.I.Sh.",
    "Otasini tel raqami",
    "Otasini ish joyi",
    "Onasini F.I.Sh.",
    "Onasini tel raqami",
    "Onasini ish joyi",
    "Telegram ID",
    "To'ldirilgan vaqt",
    "Oxirgi tahrir",
)
"""The report's columns, in the order the tutors' office asks for them."""

FULL_COLUMN_WIDTHS: tuple[int, ...] = (
    5, 28, 18, 30, 16, 24, 14, 8, 18, 20, 18, 16, 22, 22, 22, 30, 30, 20,
    28, 20, 30, 20, 22, 28, 26, 20, 34, 28, 18, 26, 28, 18, 26, 14, 20, 20,
)
_FULL_TEXT_COLUMNS = {3, 5, 9, 10, 16, 22, 26, 29, 32}  # 1-based: every phone, passport and PNFL cell

EMPLOYED_LABELS = ("Ishlamaydi", "Ishlaydi")
MARRIED_LABELS = ("Oila qurmagan", "Oila qurgan")


def full_profile_row(index: int, profile: FullProfile) -> list[Any]:
    """One report line, in :data:`FULL_HEADERS` order."""
    return [
        index,
        profile.tutor_name,
        str(profile.tutor_phone),
        profile.full_name,
        str(profile.phone),
        profile.direction,
        profile.group_name,
        course_label(profile.course),
        str(profile.passport),
        str(profile.pinfl),
        profile.birth_date,
        profile.citizenship,
        profile.region,
        profile.district,
        profile.mfy,
        str(profile.mfy_contact),
        profile.street,
        EMPLOYED_LABELS[profile.employed],
        profile.work_place,
        profile.work_position,
        profile.work_address,
        str(profile.work_phone),
        MARRIED_LABELS[profile.married],
        profile.spouse_name,
        profile.spouse_work,
        str(profile.spouse_phone),
        social_labels(profile.social_codes),
        profile.father_name,
        str(profile.father_phone),
        profile.father_work,
        profile.mother_name,
        str(profile.mother_phone),
        profile.mother_work,
        profile.telegram_id,
        profile.created_at,
        profile.edited_at,
    ]


def _fill_full_sheet(ws: Worksheet, profiles: Sequence[FullProfile]) -> None:
    ws.append(list(FULL_HEADERS))
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _BORDER

    for index, profile in enumerate(profiles, start=1):
        ws.append(full_profile_row(index, profile))
        for col_index, cell in enumerate(ws[index + 1], start=1):
            cell.border = _BORDER
            if col_index in _FULL_TEXT_COLUMNS:
                cell.number_format = "@"

    for col_index, width in enumerate(FULL_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col_index)].width = width
    ws.freeze_panes = "B2"  # the numbering column stays visible while scrolling right
    ws.auto_filter.ref = f"A1:{get_column_letter(len(FULL_HEADERS))}{max(ws.max_row, 1)}"


def build_full_workbook(sheets: Sequence[tuple[str, Sequence[FullProfile]]]) -> bytes:
    """Full-survey workbook with one sheet per ``(name, profiles)`` pair."""
    wb = Workbook()
    default = wb.active
    if default is not None:
        wb.remove(default)
    if not sheets:
        sheets = [(ALL_SHEET_NAME, [])]
    names = unique_sheet_names(name for name, _ in sheets)
    for title, (_, profiles) in zip(names, sheets, strict=True):
        _fill_full_sheet(wb.create_sheet(title=title), profiles)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_full_group_workbook(group: Group, profiles: Sequence[FullProfile]) -> bytes:
    return build_full_workbook([(group.name, profiles)])


def build_full_tutor_workbook(groups: Sequence[Group], profiles: Sequence[FullProfile]) -> bytes:
    """Sheet ``Barchasi`` with all of the tutor's profiles + one sheet per group."""
    sheets: list[tuple[str, Sequence[FullProfile]]] = [(ALL_SHEET_NAME, profiles)]
    for group in groups:
        sheets.append((group.name, [p for p in profiles if p.group_id == group.id]))
    return build_full_workbook(sheets)


def build_full_all_tutors_workbook(tutors: Sequence[Tutor], profiles: Sequence[FullProfile]) -> bytes:
    sheets: list[tuple[str, Sequence[FullProfile]]] = [(ALL_SHEET_NAME, profiles)]
    for tutor in tutors:
        sheets.append((tutor.name, [p for p in profiles if p.tutor_id == tutor.id]))
    return build_full_workbook(sheets)
