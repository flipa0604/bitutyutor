"""Integration tests through the real aiogram Dispatcher and routers with a fake Telegram session.

Self-contained harness (does not reuse tests/helpers.py) so that it independently verifies the bot:
- ``FakeSession`` subclasses ``aiogram.client.session.base.BaseSession`` and records every API call;
- the Dispatcher is built exactly as ``main.py`` does (``main.create_dispatcher``);
- ``Settings`` are produced by the real ``bot.config.load_settings`` from monkeypatched env vars.

User ids: SUPERADMIN=1001, TUTOR=2002, BOTH=3003 (superadmin AND tutor), STUDENT=4004, OTHER_TUTOR=5005.
"""

from __future__ import annotations

import asyncio
import re
from collections import Counter
from collections.abc import AsyncGenerator, AsyncIterator
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from itertools import count
from pathlib import Path
from typing import Any

import pytest
import pytest_asyncio
from aiogram import Bot, Dispatcher
from aiogram.client.default import DefaultBotProperties
from aiogram.client.session.base import BaseSession
from aiogram.enums import ParseMode
from aiogram.methods import TelegramMethod
from aiogram.methods.base import TelegramType
from aiogram.types import (
    BufferedInputFile,
    CallbackQuery,
    Chat,
    Contact,
    Document,
    InlineKeyboardMarkup,
    KeyboardButton,
    Message,
    ReplyKeyboardMarkup,
    ReplyKeyboardRemove,
    Update,
    User,
)
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from bot import texts
from bot.config import Settings, load_settings
from bot.db import Database
from bot.states import AdminTutorAdd, AdminTutorEdit, Registration, TutorGroupAdd, TutorGroupEdit
from main import create_dispatcher

SUPERADMIN = 1001
TUTOR = 2002
BOTH = 3003  # superadmin AND tutor
STUDENT = 4004
OTHER_TUTOR = 5005

SUPERADMIN_IDS = {SUPERADMIN, BOTH}

# Column headers exactly as written in SPEC §10.
SPEC_HEADERS = [
    "№",
    "F.I.SH",
    "Telefon",
    "Yo'nalish",
    "Tyutor",
    "Guruh",
    "Turar joy",
    "Manzil",
    "Otasining F.I.SH",
    "Otasining tel",
    "Onasining F.I.SH",
    "Onasining tel",
    "Telegram username",
    "Telegram ID",
    "Ro'yxatdan o'tgan vaqt",
    "Oxirgi tahrir",
]

_update_ids = count(1)
_message_ids = count(1)
_callback_ids = count(1)


# ------------------------------------------------------------------ fake session


class FakeSession(BaseSession):
    """Records ``(method name, method)`` for every request and returns plausible results."""

    def __init__(self) -> None:
        super().__init__()
        self.requests: list[tuple[str, TelegramMethod[Any]]] = []
        self._out_ids = count(100_000)

    async def close(self) -> None:
        return None

    async def stream_content(
        self,
        url: str,
        headers: dict[str, Any] | None = None,
        timeout: int = 30,
        chunk_size: int = 65536,
        raise_for_status: bool = True,
    ) -> AsyncGenerator[bytes, None]:
        raise NotImplementedError
        yield b""  # pragma: no cover - makes this an async generator as the ABC demands

    async def make_request(
        self, bot: Bot, method: TelegramMethod[TelegramType], timeout: int | None = None
    ) -> TelegramType:
        name = type(method).__name__
        self.requests.append((name, method))
        chat_id = int(getattr(method, "chat_id", 0) or 0)
        if name in {"SendMessage", "EditMessageText"}:
            return Message(  # type: ignore[return-value]
                message_id=next(self._out_ids),
                date=datetime.now(),
                chat=Chat(id=chat_id, type="private"),
                text=getattr(method, "text", None) or "x",
            )
        if name == "SendDocument":
            document = getattr(method, "document", None)
            filename = document.filename if isinstance(document, BufferedInputFile) else str(document)
            return Message(  # type: ignore[return-value]
                message_id=next(self._out_ids),
                date=datetime.now(),
                chat=Chat(id=chat_id, type="private"),
                document=Document(file_id="f", file_unique_id="fu", file_name=filename),
            )
        if name == "GetMe":
            return User(id=42, is_bot=True, first_name="bot")  # type: ignore[return-value]
        # AnswerCallbackQuery, SetMyCommands, DeleteMyCommands, DeleteMessage, EditMessageReplyMarkup, ...
        return True  # type: ignore[return-value]


# ----------------------------------------------------------------- update factories


def make_user(user_id: int, username: str | None = None, first_name: str = "User") -> User:
    return User(id=user_id, is_bot=False, first_name=first_name, username=username)


def _message(user: User, **fields: Any) -> Message:
    return Message(
        message_id=next(_message_ids),
        date=datetime.now(),
        chat=Chat(id=user.id, type="private", first_name=user.first_name),
        from_user=user,
        **fields,
    )


def text_update(user: User, text: str) -> Update:
    return Update(update_id=next(_update_ids), message=_message(user, text=text))


def contact_update(user: User, phone_number: str, contact_user_id: int | None) -> Update:
    contact = Contact(phone_number=phone_number, first_name="Contact", user_id=contact_user_id)
    return Update(update_id=next(_update_ids), message=_message(user, contact=contact))


def callback_update(user: User, data: str, message_text: str = "bot menu") -> Update:
    """A callback query whose ``message`` is a proper ``types.Message`` (so ``edit_text`` works)."""
    bot_message = Message(
        message_id=next(_message_ids),
        date=datetime.now(),
        chat=Chat(id=user.id, type="private"),
        from_user=User(id=42, is_bot=True, first_name="bot"),
        text=message_text,
    )
    callback = CallbackQuery(
        id=str(next(_callback_ids)),
        from_user=user,
        chat_instance="chat-instance",
        message=bot_message,
        data=data,
    )
    return Update(update_id=next(_update_ids), callback_query=callback)


# ------------------------------------------------------------------------ harness


@dataclass
class Harness:
    session: FakeSession
    bot: Bot
    dp: Dispatcher
    db: Database
    settings: Settings
    callbacks_fed: int = field(default=0)

    async def feed(self, *updates: Update) -> None:
        for update in updates:
            if update.callback_query is not None:
                self.callbacks_fed += 1
            await self.dp.feed_update(self.bot, update)

    async def feed_concurrently(self, *updates: Update) -> list[Any]:
        """Handle ``updates`` as concurrent tasks, the way ``start_polling`` (handle_as_tasks=True) does."""
        self.callbacks_fed += sum(1 for u in updates if u.callback_query is not None)
        return await asyncio.gather(*(self.dp.feed_update(self.bot, u) for u in updates), return_exceptions=True)

    # -- recorded requests -------------------------------------------------

    def of(self, name: str, chat_id: int | None = None) -> list[Any]:
        return [
            m
            for n, m in self.session.requests
            if n == name and (chat_id is None or int(getattr(m, "chat_id", -1)) == chat_id)
        ]

    def messages(self, chat_id: int) -> list[Any]:
        return self.of("SendMessage", chat_id)

    def texts_to(self, chat_id: int) -> list[str]:
        """Texts shown to the user in ``chat_id`` (new messages and in-place edits, in order)."""
        return [
            str(m.text)
            for n, m in self.session.requests
            if n in {"SendMessage", "EditMessageText"} and int(m.chat_id) == chat_id
        ]

    def last_text(self, chat_id: int) -> str:
        shown = self.texts_to(chat_id)
        return shown[-1] if shown else ""

    def last_message(self, chat_id: int) -> Any:
        sent = self.messages(chat_id)
        assert sent, f"nothing was sent to {chat_id}"
        return sent[-1]

    def last_shown(self, chat_id: int) -> Any:
        """Last thing displayed to the user: a new message or an in-place edit (either is acceptable)."""
        shown = [
            m
            for n, m in self.session.requests
            if n in {"SendMessage", "EditMessageText"} and int(m.chat_id) == chat_id
        ]
        assert shown, f"nothing was shown to {chat_id}"
        return shown[-1]

    def cards(self, title: str) -> Counter[int]:
        """Chat ids that received a notification card starting with ``title`` (with multiplicity)."""
        return Counter(int(m.chat_id) for m in self.of("SendMessage") if str(m.text or "").startswith(title))

    def answers(self) -> list[Any]:
        return self.of("AnswerCallbackQuery")

    def documents(self, chat_id: int | None = None) -> list[Any]:
        return self.of("SendDocument", chat_id)

    def clear(self) -> None:
        self.session.requests.clear()
        self.callbacks_fed = 0

    # -- FSM ---------------------------------------------------------------

    def ctx(self, user_id: int) -> Any:
        return self.dp.fsm.get_context(self.bot, chat_id=user_id, user_id=user_id)

    async def state_of(self, user_id: int) -> str | None:
        return await self.ctx(user_id).get_state()

    async def data_of(self, user_id: int) -> dict[str, Any]:
        return await self.ctx(user_id).get_data()


@pytest_asyncio.fixture
async def h(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> AsyncIterator[Harness]:
    # Settings the way the app builds them: from the environment (whitespace-tolerant SUPERADMIN_IDS).
    monkeypatch.setenv("BOT_TOKEN", "42:TEST")
    monkeypatch.setenv("SUPERADMIN_IDS", " 1001 , 3003 ")
    monkeypatch.setenv("DB_PATH", str(tmp_path / "flows.db"))
    settings = load_settings(env_file=tmp_path / "does-not-exist.env")
    assert settings.superadmin_ids == frozenset(SUPERADMIN_IDS)

    db = Database(settings.db_path)
    await db.init()
    session = FakeSession()
    bot = Bot("42:TEST", session=session, default=DefaultBotProperties(parse_mode=ParseMode.HTML))
    dp = create_dispatcher(db, settings)
    assert dp["db"] is db and dp["settings"] is settings
    harness = Harness(session=session, bot=bot, dp=dp, db=db, settings=settings)
    yield harness
    await db.close()
    await bot.session.close()


# ------------------------------------------------------------------ keyboard helpers


def reply_buttons(markup: Any) -> list[KeyboardButton]:
    assert isinstance(markup, ReplyKeyboardMarkup), f"expected a reply keyboard, got {type(markup).__name__}"
    return [b for row in markup.keyboard for b in row]


def reply_button_texts(markup: Any) -> list[str]:
    return [b.text for b in reply_buttons(markup)]


def inline_buttons(markup: Any) -> dict[str, str]:
    """``{callback_data: text}`` of an inline keyboard."""
    assert isinstance(markup, InlineKeyboardMarkup), f"expected an inline keyboard, got {type(markup).__name__}"
    return {str(b.callback_data): b.text for row in markup.inline_keyboard for b in row}


def load_xlsx(document: Any) -> Any:
    assert isinstance(document, BufferedInputFile)
    assert document.filename.endswith(".xlsx")
    return load_workbook(BytesIO(document.data))


def sheet_rows(ws: Worksheet) -> list[list[Any]]:
    return [list(row) for row in ws.iter_rows(values_only=True)]


def assert_all_callbacks_answered(h: Harness) -> None:
    assert len(h.answers()) == h.callbacks_fed, "every CallbackQuery must be answered exactly once"


# ----------------------------------------------------------------- registration driver


@dataclass(frozen=True)
class RegInput:
    phone: str = "+998901234567"
    full_name: str = "Aliyev Vali G'aniyevich"
    direction: str = "Dasturiy injiniring"
    residence_button: str = texts.BTN_RES_TTJ
    address: str | None = None  # required for every residence but TTJ
    father_name: str = "Aliyev G'ani Karimovich"
    father_phone: str = "+998901111111"
    mother_name: str = "Aliyeva Zulfiya Anvarovna"
    mother_phone: str = "+998902222222"


async def register(
    h: Harness,
    user: User,
    tutor_id: int,
    group_id: int,
    reg: RegInput = RegInput(),
    *,
    start_text: str = "/start",
    confirm: bool = True,
    already_registered: bool = False,
) -> None:
    """Drive the registration FSM step by step, asserting each prompt as SPEC §8 describes it.

    A student who already has a saved row lands on their own card, so the flow is reopened with the
    "register again" button instead of starting straight from the tutor picker -- which only works
    for a test user, since nobody else is allowed to overwrite a finished registration.
    """
    uid = user.id
    await h.feed(text_update(user, start_text))
    assert "Qaysi anketani" in h.last_text(uid)  # the picker comes first now
    await h.feed(callback_update(user, "sv:fill:basic"))
    if already_registered:
        assert texts.STUDENT_HOME_TITLE in h.last_text(uid)
        await h.feed(callback_update(user, "edt:again:"))
    picker = h.last_message(uid)
    assert picker.text == texts.REG_CHOOSE_TUTOR and "Tyutoringizni tanlang" in picker.text
    tutor_buttons = inline_buttons(picker.reply_markup)
    assert f"reg:tutor:{tutor_id}" in tutor_buttons, tutor_buttons
    assert await h.state_of(uid) == Registration.choose_tutor.state

    await h.feed(callback_update(user, f"reg:tutor:{tutor_id}"))
    assert h.last_text(uid) == texts.REG_CHOOSE_GROUP
    group_kb = inline_buttons(h.last_shown(uid).reply_markup)
    assert f"reg:group:{group_id}" in group_kb and "reg:back:0" in group_kb
    assert await h.state_of(uid) == Registration.choose_group.state

    await h.feed(callback_update(user, f"reg:group:{group_id}"))
    phone_prompt = h.last_message(uid)
    assert phone_prompt.text == texts.REG_ASK_PHONE
    buttons = reply_buttons(phone_prompt.reply_markup)
    assert any(b.request_contact for b in buttons) and texts.BTN_CANCEL in [b.text for b in buttons]
    assert await h.state_of(uid) == Registration.phone.state

    await h.feed(text_update(user, reg.phone))
    assert h.last_text(uid) == texts.REG_ASK_FULL_NAME
    assert reply_button_texts(h.last_message(uid).reply_markup) == [texts.BTN_CANCEL]

    await h.feed(text_update(user, reg.full_name))
    assert h.last_text(uid) == texts.REG_ASK_DIRECTION

    await h.feed(text_update(user, reg.direction))
    residence_prompt = h.last_message(uid)
    assert residence_prompt.text == texts.REG_ASK_RESIDENCE
    assert reply_button_texts(residence_prompt.reply_markup) == [
        texts.BTN_RES_TTJ,
        texts.BTN_RES_KVARTIRA,
        texts.BTN_RES_UY,
        texts.BTN_RES_QARINDOSH,
        texts.BTN_CANCEL,
    ]

    await h.feed(text_update(user, reg.residence_button))
    if reg.residence_button == texts.BTN_RES_TTJ:
        # TTJ skips the address question entirely
        assert h.last_text(uid) == texts.REG_ASK_FATHER_NAME
        assert texts.REG_ASK_ADDRESS not in h.texts_to(uid)
    else:
        assert h.last_text(uid) == texts.REG_ASK_ADDRESS
        assert reg.address is not None
        await h.feed(text_update(user, reg.address))
        assert h.last_text(uid) == texts.REG_ASK_FATHER_NAME

    await h.feed(text_update(user, reg.father_name))
    assert h.last_text(uid) == texts.REG_ASK_FATHER_PHONE
    await h.feed(text_update(user, reg.father_phone))
    assert h.last_text(uid) == texts.REG_ASK_MOTHER_NAME
    await h.feed(text_update(user, reg.mother_name))
    assert h.last_text(uid) == texts.REG_ASK_MOTHER_PHONE
    await h.feed(text_update(user, reg.mother_phone))

    preview = h.last_message(uid)
    assert preview.text.startswith(texts.REG_PREVIEW_TITLE)
    assert reg.full_name in preview.text and reg.phone in preview.text
    confirm_kb = inline_buttons(preview.reply_markup)
    assert {"reg:confirm:0", "reg:restart:0", "reg:cancel:0"} <= set(confirm_kb)
    assert await h.state_of(uid) == Registration.confirm.state
    if confirm:
        await h.feed(callback_update(user, "reg:confirm:0"))


# ================================================================== (a) TTJ path


async def test_registration_ttj_full_path(h: Harness) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    group = await h.db.add_group(tutor.id, "DI-21")
    student = make_user(STUDENT, username="vali_a")

    await register(h, student, tutor.id, group.id)

    # database row
    row = await h.db.get_student_by_telegram_id(STUDENT)
    assert row is not None
    assert row.residence == "ttj" and row.address == "TTJ"
    assert row.tutor_id == tutor.id and row.group_id == group.id
    assert row.full_name == "Aliyev Vali G'aniyevich" and row.phone == "+998901234567"
    assert row.direction == "Dasturiy injiniring"
    assert row.father_name == "Aliyev G'ani Karimovich" and row.father_phone == "+998901111111"
    assert row.mother_name == "Aliyeva Zulfiya Anvarovna" and row.mother_phone == "+998902222222"
    assert row.username == "vali_a"
    assert len(await h.db.list_students()) == 1

    # state cleared, confirmation with the keyboard removed (plain student has no roles)
    assert await h.state_of(STUDENT) is None
    saved = [m for m in h.messages(STUDENT) if m.text == texts.REG_SAVED]
    assert len(saved) == 1 and isinstance(saved[0].reply_markup, ReplyKeyboardRemove)
    assert "Ma'lumotlaringiz saqlandi" in saved[0].text

    # notifications: tutor + every superadmin exactly once, nobody else
    cards = h.cards(texts.CARD_TITLE_NEW)
    assert cards == Counter({TUTOR: 1, SUPERADMIN: 1, BOTH: 1}), cards
    assert not h.cards(texts.CARD_TITLE_UPDATE)
    card_text = next(m.text for m in h.messages(TUTOR) if m.text.startswith(texts.CARD_TITLE_NEW))
    for line in (
        "Yangi talaba ro'yxatdan o'tdi",
        "Tyutor: Karimov Aziz",
        "Guruh: DI-21",
        "F.I.SH: Aliyev Vali G'aniyevich",
        "Telefon: +998901234567",
        "Yo'nalish: Dasturiy injiniring",
        "Turar joy: TTJ",
        "Manzil: TTJ",
        "Otasi: Aliyev G'ani Karimovich",
        "Otasining tel: +998901111111",
        "Onasi: Aliyeva Zulfiya Anvarovna",
        "Onasining tel: +998902222222",
        f"Telegram: @vali_a (ID: {STUDENT})",
    ):
        assert line in card_text, line
    assert row.created_at in card_text
    # all superadmins/tutor got the very same card
    assert all(m.text == card_text for m in h.of("SendMessage") if m.text.startswith(texts.CARD_TITLE_NEW))
    # the student never receives the notification card, only the preview
    assert not any(m.text.startswith(texts.CARD_TITLE_NEW) for m in h.messages(STUDENT))

    assert_all_callbacks_answered(h)


async def test_card_omits_username_when_absent(h: Harness) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    group = await h.db.add_group(tutor.id, "DI-21")
    await register(h, make_user(STUDENT, username=None), tutor.id, group.id)
    card_text = next(m.text for m in h.messages(TUTOR) if m.text.startswith(texts.CARD_TITLE_NEW))
    assert "@" not in card_text and f"(ID: {STUDENT})" in card_text


# ================================================================== (b) kvartira / uy / qarindosh paths


@pytest.mark.parametrize(
    ("button", "code"),
    [
        (texts.BTN_RES_KVARTIRA, "kvartira"),
        (texts.BTN_RES_UY, "uy"),
        (texts.BTN_RES_QARINDOSH, "qarindosh"),
    ],
    ids=["kvartira", "uy", "qarindosh"],
)
async def test_registration_asks_address_for_non_ttj(h: Harness, button: str, code: str) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    group = await h.db.add_group(tutor.id, "DI-21")
    student = make_user(STUDENT)

    reg = RegInput(residence_button=button, address="Toshkent, Chilonzor 5, 12-uy")
    await register(h, student, tutor.id, group.id, reg)

    row = await h.db.get_student_by_telegram_id(STUDENT)
    assert row is not None
    assert row.residence == code
    assert row.address == "Toshkent, Chilonzor 5, 12-uy"
    label = {"kvartira": "Kvartira", "uy": "O'z uyi", "qarindosh": "Qarindoshinikida"}[code]
    card_text = next(m.text for m in h.messages(TUTOR) if m.text.startswith(texts.CARD_TITLE_NEW))
    assert f"Turar joy: {label}" in card_text and "Manzil: Toshkent, Chilonzor 5, 12-uy" in card_text
    assert h.cards(texts.CARD_TITLE_NEW) == Counter({TUTOR: 1, SUPERADMIN: 1, BOTH: 1})


async def test_address_validation_and_typed_residence_words(h: Harness) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    group = await h.db.add_group(tutor.id, "DI-21")
    student = make_user(STUDENT)
    await h.feed(
        text_update(student, "/start"),
        callback_update(student, "sv:fill:basic"),
        callback_update(student, f"reg:tutor:{tutor.id}"),
        callback_update(student, f"reg:group:{group.id}"),
        text_update(student, "+998901234567"),
        text_update(student, "Aliyev Vali"),
        text_update(student, "Dasturiy injiniring"),
    )
    await h.feed(text_update(student, "Marsda"))  # not a residence
    assert h.last_text(STUDENT) == texts.REG_RESIDENCE_INVALID
    assert await h.state_of(STUDENT) == Registration.residence.state

    await h.feed(text_update(student, "Kvartira"))  # the same word typed, without the emoji
    assert h.last_text(STUDENT) == texts.REG_ASK_ADDRESS
    assert (await h.data_of(STUDENT))["residence"] == "kvartira"

    await h.feed(text_update(student, "Tosh"))  # < 5 chars
    assert h.last_text(STUDENT) == texts.REG_ADDRESS_INVALID
    assert await h.state_of(STUDENT) == Registration.address.state
    await h.feed(text_update(student, "x" * 301))
    assert h.last_text(STUDENT) == texts.REG_ADDRESS_INVALID
    await h.feed(text_update(student, "Toshkent, Yunusobod 4"))
    assert h.last_text(STUDENT) == texts.REG_ASK_FATHER_NAME
    assert (await h.data_of(STUDENT))["address"] == "Toshkent, Yunusobod 4"


@pytest.mark.parametrize(
    ("typed", "code"),
    [
        ("O’zimning uyimda", "uy"),
        ("Oʻzimning uyimda", "uy"),
        ("O‘z uyi", "uy"),
        ("o'zimning uyimda", "uy"),
        ("Qarindoshimnikida", "qarindosh"),
        ("qarindoshnikida", "qarindosh"),
    ],
    ids=["U+2019", "U+02BB", "U+2018", "ascii", "qarindoshim", "qarindosh"],
)
async def test_typed_residence_accepts_phone_keyboard_apostrophes(h: Harness, typed: str, code: str) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    group = await h.db.add_group(tutor.id, "DI-21")
    student = make_user(STUDENT)
    await h.feed(
        text_update(student, "/start"),
        callback_update(student, "sv:fill:basic"),
        callback_update(student, f"reg:tutor:{tutor.id}"),
        callback_update(student, f"reg:group:{group.id}"),
        text_update(student, "+998901234567"),
        text_update(student, "Aliyev Vali"),
        text_update(student, "Dasturiy injiniring"),
    )
    await h.feed(text_update(student, typed))
    assert h.last_text(STUDENT) == texts.REG_ASK_ADDRESS
    assert await h.state_of(STUDENT) == Registration.address.state
    assert (await h.data_of(STUDENT))["residence"] == code


# ================================================================== (c) tutor == superadmin


async def test_tutor_who_is_superadmin_gets_exactly_one_notification(h: Harness) -> None:
    tutor = await h.db.add_tutor("Admin Tyutor", BOTH)
    group = await h.db.add_group(tutor.id, "AI-22")

    await register(h, make_user(STUDENT), tutor.id, group.id)

    cards = h.cards(texts.CARD_TITLE_NEW)
    assert cards[BOTH] == 1, cards
    assert cards == Counter({BOTH: 1, SUPERADMIN: 1})
    assert TUTOR not in cards and OTHER_TUTOR not in cards
    assert (await h.db.get_student_by_telegram_id(STUDENT)) is not None


# ================================================================== (d) phone validation


async def _go_to_phone_step(h: Harness, student: User) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    group = await h.db.add_group(tutor.id, "DI-21")
    await h.feed(
        text_update(student, "/start"),
        callback_update(student, "sv:fill:basic"),
        callback_update(student, f"reg:tutor:{tutor.id}"),
        callback_update(student, f"reg:group:{group.id}"),
    )
    assert h.last_text(student.id) == texts.REG_ASK_PHONE
    assert await h.state_of(student.id) == Registration.phone.state


async def test_phone_text_validation(h: Harness) -> None:
    student = make_user(STUDENT)
    await _go_to_phone_step(h, student)

    for bad in ("901234567", "+99890123456", "998901234567", "+9989012345678", "+998 90 123 45 67", "hello"):
        h.clear()
        await h.feed(text_update(student, bad))
        assert h.last_text(STUDENT) == texts.REG_PHONE_INVALID, bad
        assert "Raqam noto'g'ri" in h.last_text(STUDENT) and "+998901234567" in h.last_text(STUDENT)
        assert await h.state_of(STUDENT) == Registration.phone.state, bad
        assert "phone" not in await h.data_of(STUDENT)
        # the contact keyboard is offered again
        assert any(b.request_contact for b in reply_buttons(h.last_message(STUDENT).reply_markup))

    await h.feed(text_update(student, "  +998901234567  "))  # valid after strip()
    assert h.last_text(STUDENT) == texts.REG_ASK_FULL_NAME
    assert await h.state_of(STUDENT) == Registration.full_name.state
    assert (await h.data_of(STUDENT))["phone"] == "+998901234567"


async def test_phone_contact_validation(h: Harness) -> None:
    student = make_user(STUDENT)
    await _go_to_phone_step(h, student)

    # someone else's contact is rejected and the state is kept
    await h.feed(contact_update(student, "998901234567", OTHER_TUTOR))
    assert h.last_text(STUDENT) == texts.REG_CONTACT_NOT_OWN
    assert "Faqat o'zingizning raqamingizni yuboring" in h.last_text(STUDENT)
    assert await h.state_of(STUDENT) == Registration.phone.state
    assert "phone" not in await h.data_of(STUDENT)

    # a contact that is not linked to any Telegram account is not "own" either
    await h.feed(contact_update(student, "998901234567", None))
    assert h.last_text(STUDENT) == texts.REG_CONTACT_NOT_OWN
    assert await h.state_of(STUDENT) == Registration.phone.state

    # own contact without '+' is normalised
    await h.feed(contact_update(student, "998901234567", STUDENT))
    assert h.last_text(STUDENT) == texts.REG_ASK_FULL_NAME
    assert await h.state_of(STUDENT) == Registration.full_name.state
    assert (await h.data_of(STUDENT))["phone"] == "+998901234567"


async def test_phone_from_own_contact_ends_up_in_db(h: Harness) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    group = await h.db.add_group(tutor.id, "DI-21")
    student = make_user(STUDENT)
    await h.feed(
        text_update(student, "/start"),
        callback_update(student, "sv:fill:basic"),
        callback_update(student, f"reg:tutor:{tutor.id}"),
        callback_update(student, f"reg:group:{group.id}"),
        contact_update(student, "998 90 765-43-21", STUDENT),
        text_update(student, "Aliyev Vali"),
        text_update(student, "Dasturiy injiniring"),
        text_update(student, texts.BTN_RES_TTJ),
        text_update(student, "Aliyev G'ani"),
        text_update(student, "+998901111111"),
        text_update(student, "Aliyeva Zulfiya"),
        text_update(student, "+998902222222"),
        callback_update(student, "reg:confirm:0"),
    )
    row = await h.db.get_student_by_telegram_id(STUDENT)
    assert row is not None and row.phone == "+998907654321"


async def test_name_and_parent_phone_validation(h: Harness) -> None:
    student = make_user(STUDENT)
    await _go_to_phone_step(h, student)
    await h.feed(text_update(student, "+998901234567"))

    for bad_name in ("Vali", "Aliyev Vali 2", "A B" * 60, " ", "Al"):
        await h.feed(text_update(student, bad_name))
        assert h.last_text(STUDENT) == texts.REG_NAME_INVALID, bad_name
        assert await h.state_of(STUDENT) == Registration.full_name.state
    await h.feed(text_update(student, "  Aliyev   Vali  "))
    assert (await h.data_of(STUDENT))["full_name"] == "Aliyev Vali"

    await h.feed(text_update(student, "D"))
    assert h.last_text(STUDENT) == texts.REG_DIRECTION_INVALID
    await h.feed(text_update(student, "Dasturiy injiniring"), text_update(student, texts.BTN_RES_TTJ))

    await h.feed(text_update(student, "Aliyev 1"))
    assert h.last_text(STUDENT) == texts.REG_NAME_INVALID
    await h.feed(text_update(student, "Aliyev G'ani"))
    assert h.last_text(STUDENT) == texts.REG_ASK_FATHER_PHONE

    await h.feed(text_update(student, "998901111111"))
    assert h.last_text(STUDENT) == texts.REG_PARENT_PHONE_INVALID
    assert await h.state_of(STUDENT) == Registration.father_phone.state
    await h.feed(text_update(student, "+998901111111"))
    assert h.last_text(STUDENT) == texts.REG_ASK_MOTHER_NAME
    await h.feed(text_update(student, "Aliyeva Zulfiya"), text_update(student, "+99890222222"))
    assert h.last_text(STUDENT) == texts.REG_PARENT_PHONE_INVALID
    assert await h.state_of(STUDENT) == Registration.mother_phone.state
    await h.feed(text_update(student, "+998902222222"))
    assert h.last_text(STUDENT).startswith(texts.REG_PREVIEW_TITLE)


# ================================================================== (e) /add_tutor and /admin guard


async def test_superadmin_add_tutor_flow(h: Harness) -> None:
    admin = make_user(SUPERADMIN)

    await h.feed(text_update(admin, "/add_tutor"))
    assert h.last_text(SUPERADMIN) == texts.ASK_TUTOR_NAME
    assert await h.state_of(SUPERADMIN) == AdminTutorAdd.name.state

    await h.feed(text_update(admin, "x" * 101))
    assert h.last_text(SUPERADMIN) == texts.TUTOR_NAME_INVALID
    await h.feed(text_update(admin, "Karimov Aziz"))
    assert h.last_text(SUPERADMIN) == texts.ASK_TUTOR_TG
    assert "@userinfobot" in h.last_text(SUPERADMIN)
    assert await h.state_of(SUPERADMIN) == AdminTutorAdd.telegram_id.state

    # "²"/"٣٤" pass str.isdigit(); 20 digits overflow SQLite's INTEGER — all must be answered, not crash
    for bad in ("abc", "-5", "0", "12.5", "99999999999999999999", str(2**63), "²", "٣٤"):
        await h.feed(text_update(admin, bad))
        assert h.last_text(SUPERADMIN) == texts.TUTOR_TG_INVALID, bad
        assert await h.state_of(SUPERADMIN) == AdminTutorAdd.telegram_id.state

    await h.feed(text_update(admin, str(TUTOR)))
    confirm = h.last_message(SUPERADMIN)
    assert "Karimov Aziz" in confirm.text and str(TUTOR) in confirm.text
    kb = inline_buttons(confirm.reply_markup)
    assert kb.get("adm:save:0") == texts.BTN_SAVE and kb.get("adm:cancel:0") == texts.BTN_CANCEL
    assert await h.state_of(SUPERADMIN) == AdminTutorAdd.confirm.state
    assert await h.db.get_tutor_by_telegram_id(TUTOR) is None  # nothing saved before confirmation

    await h.feed(callback_update(admin, "adm:save:0"))
    tutor = await h.db.get_tutor_by_telegram_id(TUTOR)
    assert tutor is not None and tutor.name == "Karimov Aziz"
    assert len(await h.db.list_tutors()) == 1
    assert texts.TUTOR_SAVED in h.texts_to(SUPERADMIN)
    assert await h.state_of(SUPERADMIN) is None
    # the new tutor's command menu is refreshed (BotCommandScopeChat for the tutor)
    scopes = [getattr(m.scope, "chat_id", None) for m in h.of("SetMyCommands")]
    assert TUTOR in scopes
    tutor_cmds = next(m for m in h.of("SetMyCommands") if getattr(m.scope, "chat_id", None) == TUTOR)
    assert {"tutor", "groups", "add_group", "edit_group", "delete_group", "excel"} <= {
        c.command for c in tutor_cmds.commands
    }
    assert "admin" not in {c.command for c in tutor_cmds.commands}
    assert_all_callbacks_answered(h)

    # the new tutor gets access immediately (live DB lookup)
    tutor_user = make_user(TUTOR)
    await h.feed(text_update(tutor_user, "/tutor"))
    assert "Tyutor panel" in h.last_text(TUTOR) and "Karimov Aziz" in h.last_text(TUTOR)

    # duplicate telegram id is refused and the admin stays on the ID question
    await h.feed(text_update(admin, "/add_tutor"), text_update(admin, "Boshqa Odam"), text_update(admin, str(TUTOR)))
    assert h.last_text(SUPERADMIN) == texts.TUTOR_TG_DUPLICATE
    assert await h.state_of(SUPERADMIN) == AdminTutorAdd.telegram_id.state
    await h.feed(text_update(admin, "/cancel"))
    assert len(await h.db.list_tutors()) == 1

    # /tutors lists the tutor as "{name} (id: {telegram_id})"
    await h.feed(text_update(admin, "/tutors"))
    assert f"Karimov Aziz (id: {TUTOR})" in inline_buttons(h.last_message(SUPERADMIN).reply_markup).values()


async def test_add_tutor_confirm_cancel_saves_nothing(h: Harness) -> None:
    admin = make_user(SUPERADMIN)
    await h.feed(text_update(admin, "/add_tutor"), text_update(admin, "Karimov Aziz"), text_update(admin, str(TUTOR)))
    await h.feed(text_update(admin, "salom"))  # text in the confirm state
    assert h.last_text(SUPERADMIN) == texts.USE_BUTTONS
    await h.feed(callback_update(admin, "adm:cancel:0"))
    assert h.last_text(SUPERADMIN) == texts.CANCELLED
    assert await h.state_of(SUPERADMIN) is None
    assert await h.db.list_tutors() == []


async def test_non_superadmin_is_refused_admin_commands(h: Harness) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    for user in (make_user(TUTOR), make_user(STUDENT), make_user(OTHER_TUTOR)):
        h.clear()
        for command in ("/admin", "/add_tutor", "/tutors", "/edit_tutor", "/delete_tutor"):
            await h.feed(text_update(user, command))
            assert h.last_text(user.id) == texts.ADMIN_ONLY, (user.id, command)
            assert await h.state_of(user.id) is None
        # admin inline buttons are dead for non-superadmins
        await h.feed(text_update(user, texts.BTN_ADMIN_PANEL))
        assert "Admin panel" not in h.last_text(user.id)
        await h.feed(callback_update(user, "adm:add:0"), callback_update(user, f"adm:delete:{tutor.id}"))
        await h.feed(callback_update(user, f"adm:confirm_delete:{tutor.id}"))
        assert len(h.answers()) == 3 and all(a.text == texts.STALE_BUTTON for a in h.answers())
        assert await h.state_of(user.id) is None
    assert await h.db.get_tutor(tutor.id) is not None

    # a superadmin (who is also a tutor) is allowed
    await h.feed(text_update(make_user(BOTH), "/admin"))
    assert "Admin panel" in h.last_text(BOTH)
    kb = inline_buttons(h.last_message(BOTH).reply_markup)
    assert {texts.BTN_ADMIN_TUTORS, texts.BTN_ADMIN_ADD_TUTOR, texts.BTN_ADMIN_EXCEL_ALL} <= set(kb.values())


# ================================================================== (f) /add_group + Excel exports


async def _seed_tutor_with_students(h: Harness) -> tuple[Any, Any, Any]:
    """Tutor TUTOR with groups DI-21 (2 students: one TTJ, one kvartira) and DI-22 (1 TTJ student)."""
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    g1 = await h.db.add_group(tutor.id, "DI-21")
    g2 = await h.db.add_group(tutor.id, "DI-22")
    await register(h, make_user(STUDENT, username="vali"), tutor.id, g1.id, RegInput(full_name="Zokirov Vali"))
    await register(
        h,
        make_user(4005, username="olim"),
        tutor.id,
        g1.id,
        RegInput(
            full_name="Aliyev Olim",
            phone="+998907777777",
            residence_button=texts.BTN_RES_KVARTIRA,
            address="Toshkent, Chilonzor 5",
        ),
    )
    await register(h, make_user(4006), tutor.id, g2.id, RegInput(full_name="Boboev Sardor", phone="+998908888888"))
    h.clear()
    return tutor, g1, g2


async def test_tutor_add_group_and_export_group_excel(h: Harness) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    tutor_user = make_user(TUTOR)

    await h.feed(text_update(tutor_user, "/add_group"))
    assert h.last_text(TUTOR) == texts.ASK_GROUP_NAME
    await h.feed(text_update(tutor_user, "x" * 65))
    assert h.last_text(TUTOR) == texts.GROUP_NAME_INVALID
    await h.feed(text_update(tutor_user, "DI-21"))
    groups = await h.db.list_groups(tutor.id)
    assert [g.name for g in groups] == ["DI-21"]
    assert texts.GROUP_SAVED in h.texts_to(TUTOR)
    assert await h.state_of(TUTOR) is None
    group = groups[0]

    # duplicate group name for the same tutor
    await h.feed(text_update(tutor_user, "/add_group"), text_update(tutor_user, "di-21".upper()))
    assert h.last_text(TUTOR) == texts.GROUP_DUPLICATE
    await h.feed(text_update(tutor_user, texts.BTN_CANCEL))
    assert len(await h.db.list_groups(tutor.id)) == 1

    # empty group -> no file
    await h.feed(text_update(tutor_user, "/excel"))
    menu = h.last_message(TUTOR)
    assert menu.text == texts.EXCEL_MENU
    menu_kb = inline_buttons(menu.reply_markup)
    assert {texts.BTN_EXCEL_BY_GROUP, texts.BTN_EXCEL_BY_RESIDENCE, texts.BTN_EXCEL_ALL_GROUPS} <= set(menu_kb.values())
    by_group = next(d for d, t in menu_kb.items() if t == texts.BTN_EXCEL_BY_GROUP)
    await h.feed(callback_update(tutor_user, by_group))
    pick_kb = inline_buttons(h.last_shown(TUTOR).reply_markup)
    group_button = next(d for d, t in pick_kb.items() if t.startswith("DI-21"))
    assert "📋 0 ta" in pick_kb[group_button]
    await h.feed(callback_update(tutor_user, group_button))
    assert h.last_text(TUTOR) == texts.NO_DATA and not h.documents()

    # register two students into the group, then export
    await register(h, make_user(STUDENT, username="vali"), tutor.id, group.id, RegInput(full_name="Zokirov Vali"))
    await register(
        h,
        make_user(4005),
        tutor.id,
        group.id,
        RegInput(
            full_name="Aliyev Olim",
            phone="+998907777777",
            residence_button=texts.BTN_RES_KVARTIRA,
            address="Toshkent, Chilonzor 5",
        ),
    )
    h.clear()
    await h.feed(callback_update(tutor_user, group_button))
    docs = h.documents()
    assert len(docs) == 1 and int(docs[0].chat_id) == TUTOR
    document = docs[0].document
    assert isinstance(document, BufferedInputFile)
    assert re.fullmatch(r"Karimov_Aziz_DI-21_\d{4}-\d{2}-\d{2}\.xlsx", document.filename), document.filename

    wb = load_xlsx(document)
    assert wb.sheetnames == ["DI-21"]
    ws = wb["DI-21"]
    rows = sheet_rows(ws)
    assert rows[0] == SPEC_HEADERS
    assert len(rows) == 3  # header + 2 students
    # sorted by name inside the group
    assert [r[1] for r in rows[1:]] == ["Aliyev Olim", "Zokirov Vali"]
    olim, vali = rows[1], rows[2]
    assert olim[0] == 1 and vali[0] == 2
    assert vali[2] == "+998901234567" and isinstance(vali[2], str)
    assert vali[4] == "Karimov Aziz" and vali[5] == "DI-21"
    assert vali[6] == "TTJ" and vali[7] == "TTJ"
    assert olim[6] == "Kvartira" and olim[7] == "Toshkent, Chilonzor 5"
    assert vali[12] == "@vali" and olim[12] is None
    assert vali[13] == STUDENT and olim[13] == 4005
    assert vali[9] == "+998901111111" and vali[11] == "+998902222222"
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref
    assert all(c.font.bold for c in ws[1])
    assert_all_callbacks_answered(h)


async def test_tutor_export_by_residence(h: Harness) -> None:
    tutor, g1, g2 = await _seed_tutor_with_students(h)
    tutor_user = make_user(TUTOR)

    await h.feed(text_update(tutor_user, "/excel"))
    menu_kb = inline_buttons(h.last_message(TUTOR).reply_markup)
    by_res = next(d for d, t in menu_kb.items() if t == texts.BTN_EXCEL_BY_RESIDENCE)
    await h.feed(callback_update(tutor_user, by_res))
    res_kb = inline_buttons(h.last_shown(TUTOR).reply_markup)
    assert {texts.BTN_RES_TTJ, texts.BTN_RES_KVARTIRA, texts.BTN_RES_QARINDOSH} <= set(res_kb.values())
    ttj_button = next(d for d, t in res_kb.items() if t == texts.BTN_RES_TTJ)
    kv_button = next(d for d, t in res_kb.items() if t == texts.BTN_RES_KVARTIRA)
    uy_button = next(d for d, t in res_kb.items() if "uyi" in t.lower())
    rel_button = next(d for d, t in res_kb.items() if t == texts.BTN_RES_QARINDOSH)

    await h.feed(callback_update(tutor_user, ttj_button))
    docs = h.documents(TUTOR)
    assert len(docs) == 1
    assert re.fullmatch(r"Karimov_Aziz_TTJ_\d{4}-\d{2}-\d{2}\.xlsx", docs[0].document.filename)
    wb = load_xlsx(docs[0].document)
    assert len(wb.sheetnames) == 1
    rows = sheet_rows(wb.worksheets[0])
    assert rows[0] == SPEC_HEADERS
    # both TTJ students across both groups, sorted by group then name; kvartira student excluded
    assert [(r[5], r[1], r[6]) for r in rows[1:]] == [("DI-21", "Zokirov Vali", "TTJ"), ("DI-22", "Boboev Sardor", "TTJ")]

    await h.feed(callback_update(tutor_user, kv_button))
    docs = h.documents(TUTOR)
    assert len(docs) == 2
    rows = sheet_rows(load_xlsx(docs[1].document).worksheets[0])
    assert [(r[1], r[6], r[7]) for r in rows[1:]] == [("Aliyev Olim", "Kvartira", "Toshkent, Chilonzor 5")]

    # nobody lives at home -> message instead of an empty file
    await h.feed(callback_update(tutor_user, uy_button))
    assert h.last_text(TUTOR) == texts.NO_DATA and len(h.documents(TUTOR)) == 2

    # a student staying with relatives: absent from the sheets above, gets a sheet of their own
    await register(
        h,
        make_user(4007),
        tutor.id,
        g2.id,
        RegInput(
            full_name="Rahimov Jasur",
            phone="+998909999999",
            residence_button=texts.BTN_RES_QARINDOSH,
            address="Samarqand, Registon 3",
        ),
    )
    await h.feed(callback_update(tutor_user, rel_button))
    docs = h.documents(TUTOR)
    assert len(docs) == 3
    assert re.fullmatch(r"Karimov_Aziz_Qarindoshinikida_\d{4}-\d{2}-\d{2}\.xlsx", docs[2].document.filename)
    wb = load_xlsx(docs[2].document)
    assert wb.sheetnames == ["Qarindoshinikida"]
    rows = sheet_rows(wb.worksheets[0])
    assert [(r[5], r[1], r[6], r[7]) for r in rows[1:]] == [
        ("DI-22", "Rahimov Jasur", "Qarindoshinikida", "Samarqand, Registon 3")
    ]

    # an invalid residence code in callback data is rejected, not exported
    await h.feed(callback_update(tutor_user, "tut:excel_res_pick:0:hotel"))
    assert len(h.documents(TUTOR)) == 3
    assert_all_callbacks_answered(h)


async def test_tutor_export_all_groups(h: Harness) -> None:
    tutor, g1, g2 = await _seed_tutor_with_students(h)
    empty = await h.db.add_group(tutor.id, "Bo'sh")
    tutor_user = make_user(TUTOR)

    await h.feed(text_update(tutor_user, "/excel"))
    menu_kb = inline_buttons(h.last_message(TUTOR).reply_markup)
    all_button = next(d for d, t in menu_kb.items() if t == texts.BTN_EXCEL_ALL_GROUPS)
    await h.feed(callback_update(tutor_user, all_button))

    docs = h.documents(TUTOR)
    assert len(docs) == 1
    assert re.fullmatch(r"Karimov_Aziz_barcha_guruhlar_\d{4}-\d{2}-\d{2}\.xlsx", docs[0].document.filename)
    wb = load_xlsx(docs[0].document)
    assert wb.sheetnames[0] == "Barchasi"
    assert set(wb.sheetnames) == {"Barchasi", "DI-21", "DI-22", "Bo'sh"}
    assert sheet_rows(wb["Barchasi"])[0] == SPEC_HEADERS
    assert len(sheet_rows(wb["Barchasi"])) == 4
    assert len(sheet_rows(wb["DI-21"])) == 3
    assert len(sheet_rows(wb["DI-22"])) == 2
    assert len(sheet_rows(wb["Bo'sh"])) == 1
    assert {r[5] for r in sheet_rows(wb["DI-22"])[1:]} == {"DI-22"}
    assert docs[0].caption and "3 ta" in docs[0].caption


async def test_superadmin_export_all_tutors(h: Harness) -> None:
    tutor, g1, g2 = await _seed_tutor_with_students(h)
    other = await h.db.add_tutor("Boshqa Tyutor", OTHER_TUTOR)
    og = await h.db.add_group(other.id, "DI-21")  # same group name at another tutor
    await register(h, make_user(4007), other.id, og.id, RegInput(full_name="Karimov Bek", phone="+998909999999"))
    h.clear()

    admin = make_user(SUPERADMIN)
    await h.feed(text_update(admin, "/admin"))
    kb = inline_buttons(h.last_message(SUPERADMIN).reply_markup)
    all_button = next(d for d, t in kb.items() if t == texts.BTN_ADMIN_EXCEL_ALL)
    await h.feed(callback_update(admin, all_button))
    docs = h.documents(SUPERADMIN)
    assert len(docs) == 1 and docs[0].document.filename.endswith(".xlsx")
    wb = load_xlsx(docs[0].document)
    assert wb.sheetnames == ["Barchasi", "Boshqa Tyutor", "Karimov Aziz"]
    assert len(sheet_rows(wb["Barchasi"])) == 5
    assert len(sheet_rows(wb["Karimov Aziz"])) == 4
    assert len(sheet_rows(wb["Boshqa Tyutor"])) == 2
    tyutor_col = {r[4] for r in sheet_rows(wb["Barchasi"])[1:]}
    assert tyutor_col == {"Karimov Aziz", "Boshqa Tyutor"}


# ================================================================== (g) foreign group isolation


async def test_tutor_cannot_touch_foreign_group(h: Harness) -> None:
    owner = await h.db.add_tutor("Karimov Aziz", TUTOR)
    intruder_tutor = await h.db.add_tutor("Boshqa Tyutor", OTHER_TUTOR)
    group = await h.db.add_group(owner.id, "DI-21")
    await register(h, make_user(STUDENT), owner.id, group.id)
    h.clear()

    intruder = make_user(OTHER_TUTOR)
    for data in (
        f"tut:excel_group:{group.id}:",
        f"tut:view:{group.id}:",
        f"tut:rename:{group.id}:",
        f"tut:delete:{group.id}:",
        f"tut:confirm_delete:{group.id}:",
    ):
        h.clear()
        await h.feed(callback_update(intruder, data))
        answers = h.answers()
        assert len(answers) == 1, data
        assert answers[0].text == texts.NO_PERMISSION and "Ruxsat yo'q" in answers[0].text, data
        assert not h.documents(), data
        assert not h.messages(OTHER_TUTOR) and not h.of("EditMessageText"), data
        assert await h.state_of(OTHER_TUTOR) is None, data

    # a rename attempt after the refused callback must not rename anything
    await h.feed(text_update(intruder, "HACKED"))
    assert h.last_text(OTHER_TUTOR) == texts.UNKNOWN

    # defence in depth: even with a forged FSM state pointing at the foreign group, rename is refused
    ctx = h.ctx(OTHER_TUTOR)
    await ctx.set_state(TutorGroupEdit.name)
    await ctx.update_data(group_id=group.id)
    await h.feed(text_update(intruder, "HACKED"))
    assert h.last_text(OTHER_TUTOR) == texts.GROUP_NOT_FOUND
    assert await h.state_of(OTHER_TUTOR) is None

    # database untouched
    still = await h.db.get_group(group.id)
    assert still is not None and still.name == "DI-21" and still.tutor_id == owner.id
    assert (await h.db.get_student_by_telegram_id(STUDENT)) is not None
    assert await h.db.list_groups(intruder_tutor.id) == []

    # the intruder sees only their own (empty) group list
    await h.feed(text_update(intruder, "/groups"))
    assert "DI-21" not in str(h.last_message(OTHER_TUTOR).reply_markup)

    # a user without the tutor role gets refused outright
    stranger = make_user(STUDENT)
    await h.feed(text_update(stranger, "/tutor"), text_update(stranger, "/excel"))
    assert h.texts_to(STUDENT)[-2:] == [texts.TUTOR_ONLY, texts.TUTOR_ONLY]
    h.clear()
    await h.feed(callback_update(stranger, f"tut:excel_group:{group.id}:"))
    assert h.answers()[-1].text == texts.STALE_BUTTON and not h.documents()

    # a superadmin who is NOT a tutor cannot open the tutor panel either (roles are not implied)
    await h.feed(text_update(make_user(SUPERADMIN), "/tutor"))
    assert h.last_text(SUPERADMIN) == texts.TUTOR_ONLY


# ================================================================== (h) /cancel mid-registration


async def test_cancel_mid_registration_clears_state(h: Harness) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    group = await h.db.add_group(tutor.id, "DI-21")
    student = make_user(STUDENT)

    await h.feed(
        text_update(student, "/start"),
        callback_update(student, "sv:fill:basic"),
        callback_update(student, f"reg:tutor:{tutor.id}"),
        callback_update(student, f"reg:group:{group.id}"),
        text_update(student, "+998901234567"),
    )
    assert await h.state_of(STUDENT) == Registration.full_name.state
    assert (await h.data_of(STUDENT))["phone"] == "+998901234567"

    await h.feed(text_update(student, "/cancel"))
    cancelled = h.last_message(STUDENT)
    assert cancelled.text == texts.CANCELLED_STUDENT and "/start" in cancelled.text
    assert isinstance(cancelled.reply_markup, ReplyKeyboardRemove)
    assert await h.state_of(STUDENT) is None
    assert await h.data_of(STUDENT) == {}

    # the next name-like text is no longer consumed by the FSM
    await h.feed(text_update(student, "Aliyev Vali"))
    assert h.last_text(STUDENT) == texts.UNKNOWN
    assert await h.state_of(STUDENT) is None
    assert await h.db.get_student_by_telegram_id(STUDENT) is None

    # cancel button in a later step behaves the same
    await h.feed(
        text_update(student, "/start"),
        callback_update(student, "sv:fill:basic"),
        callback_update(student, f"reg:tutor:{tutor.id}"),
        callback_update(student, f"reg:group:{group.id}"),
        text_update(student, "+998901234567"),
        text_update(student, "Aliyev Vali"),
        text_update(student, "Dasturiy injiniring"),
    )
    assert await h.state_of(STUDENT) == Registration.residence.state
    await h.feed(text_update(student, texts.BTN_CANCEL))
    assert h.last_text(STUDENT) == texts.CANCELLED_STUDENT
    assert await h.state_of(STUDENT) is None and await h.data_of(STUDENT) == {}

    # inline cancel on the confirm card
    await register(h, student, tutor.id, group.id, confirm=False)
    await h.feed(callback_update(student, "reg:cancel:0"))
    assert h.last_text(STUDENT) == texts.CANCELLED_STUDENT
    assert await h.state_of(STUDENT) is None
    assert await h.db.get_student_by_telegram_id(STUDENT) is None
    assert not h.cards(texts.CARD_TITLE_NEW)


async def test_cancel_for_role_user_returns_to_main_menu(h: Harness) -> None:
    tutor = await h.db.add_tutor("Admin Tyutor", BOTH)
    group = await h.db.add_group(tutor.id, "AI-22")
    user = make_user(BOTH)
    await h.feed(
        text_update(user, texts.BTN_REGISTER),
        callback_update(user, "sv:fill:basic"),
        callback_update(user, f"reg:tutor:{tutor.id}"),
        callback_update(user, f"reg:group:{group.id}"),
    )
    assert await h.state_of(BOTH) == Registration.phone.state
    await h.feed(text_update(user, "/cancel"))
    cancelled = h.last_message(BOTH)
    assert cancelled.text == texts.CANCELLED
    assert set(reply_button_texts(cancelled.reply_markup)) == {
        texts.BTN_ADMIN_PANEL,
        texts.BTN_TUTOR_PANEL,
        texts.BTN_REGISTER,
    }
    assert await h.state_of(BOTH) is None


# ================================================================== (i) /start for BOTH


async def test_start_for_both_roles_shows_additive_menu(h: Harness) -> None:
    await h.db.add_tutor("Admin Tyutor", BOTH)
    user = make_user(BOTH, first_name="Boss")

    await h.feed(text_update(user, "/start"))
    greeting = h.last_message(BOTH)
    assert "Boss" in greeting.text
    buttons = reply_button_texts(greeting.reply_markup)
    assert texts.BTN_ADMIN_PANEL in buttons and texts.BTN_TUTOR_PANEL in buttons and texts.BTN_REGISTER in buttons
    assert await h.state_of(BOTH) is None
    assert texts.REG_CHOOSE_TUTOR not in h.texts_to(BOTH)

    # both panels are reachable through those buttons
    await h.feed(text_update(user, texts.BTN_ADMIN_PANEL))
    assert "Admin panel" in h.last_text(BOTH)
    await h.feed(text_update(user, texts.BTN_TUTOR_PANEL))
    assert "Tyutor panel" in h.last_text(BOTH) and "Admin Tyutor" in h.last_text(BOTH)
    await h.feed(text_update(user, "/help"))
    assert "/add_tutor" in h.last_text(BOTH) and "/add_group" in h.last_text(BOTH)

    # ... and the register button opens the survey picker for a role user too
    await h.feed(text_update(user, texts.BTN_REGISTER))
    assert "Qaysi anketani" in h.last_text(BOTH)
    await h.feed(callback_update(user, "sv:fill:basic"))
    assert h.last_text(BOTH) == texts.REG_CHOOSE_TUTOR


async def test_start_menus_for_single_roles_and_students(h: Harness) -> None:
    await h.db.add_tutor("Karimov Aziz", TUTOR)

    await h.feed(text_update(make_user(SUPERADMIN), "/start"))
    buttons = reply_button_texts(h.last_message(SUPERADMIN).reply_markup)
    assert buttons == [texts.BTN_ADMIN_PANEL, texts.BTN_REGISTER]

    await h.feed(text_update(make_user(TUTOR), "/start"))
    buttons = reply_button_texts(h.last_message(TUTOR).reply_markup)
    assert buttons == [texts.BTN_TUTOR_PANEL, texts.BTN_REGISTER]

    # plain student is asked which questionnaire first, then goes into it
    await h.feed(text_update(make_user(STUDENT), "/start"))
    assert "Qaysi anketani" in h.last_text(STUDENT)
    assert inline_buttons(h.last_message(STUDENT).reply_markup) == {
        "sv:fill:basic": f"🕗 {texts.BTN_SURVEY_BASIC}",
        "sv:fill:full": f"🕗 {texts.BTN_SURVEY_FULL}",
    }
    await h.feed(callback_update(make_user(STUDENT), "sv:fill:basic"))
    assert h.last_text(STUDENT) == texts.REG_CHOOSE_TUTOR
    assert await h.state_of(STUDENT) == Registration.choose_tutor.state

    # unknown text without a state
    await h.feed(text_update(make_user(STUDENT), "/cancel"), text_update(make_user(STUDENT), "salom"))
    assert h.last_text(STUDENT) == texts.UNKNOWN and "/start" in h.last_text(STUDENT)


async def test_start_without_tutors_stops(h: Harness) -> None:
    student = make_user(STUDENT)
    await h.feed(text_update(student, "/start"), callback_update(student, "sv:fill:basic"))
    assert h.last_text(STUDENT) == texts.REG_NO_TUTORS
    assert await h.state_of(STUDENT) is None
    await h.feed(callback_update(student, "sv:fill:full"))  # the full survey needs tutors just as much
    assert h.last_text(STUDENT) == texts.REG_NO_TUTORS
    assert await h.state_of(STUDENT) is None


# ================================================================== (j) re-registration = update


async def test_reregistration_updates_instead_of_duplicating(h: Harness) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    g1 = await h.db.add_group(tutor.id, "DI-21")
    g2 = await h.db.add_group(tutor.id, "DI-22")
    student = make_user(STUDENT, username="vali")

    await register(h, student, tutor.id, g1.id)
    first = await h.db.get_student_by_telegram_id(STUDENT)
    assert first is not None
    assert h.cards(texts.CARD_TITLE_NEW) == Counter({TUTOR: 1, SUPERADMIN: 1, BOTH: 1})
    h.clear()

    await h.db.add_test_user(STUDENT)  # only testers may run the flow a second time
    await register(
        h,
        student,
        tutor.id,
        g2.id,
        RegInput(
            phone="+998905555555",
            full_name="Aliyev Vali",
            residence_button=texts.BTN_RES_KVARTIRA,
            address="Toshkent, Yakkasaroy 7",
        ),
        already_registered=True,
    )
    rows = await h.db.list_students()
    assert len(rows) == 1
    second = rows[0]
    assert second.id == first.id and second.telegram_id == STUDENT
    assert second.group_id == g2.id and second.group_name == "DI-22"
    assert second.phone == "+998905555555" and second.full_name == "Aliyev Vali"
    assert second.residence == "kvartira" and second.address == "Toshkent, Yakkasaroy 7"
    assert second.created_at == first.created_at
    assert second.updated_at >= first.updated_at

    # notification labelled as an update, again exactly once per recipient, no "new" card
    assert h.cards(texts.CARD_TITLE_UPDATE) == Counter({TUTOR: 1, SUPERADMIN: 1, BOTH: 1})
    assert not h.cards(texts.CARD_TITLE_NEW)
    card_text = next(m.text for m in h.messages(SUPERADMIN) if m.text.startswith(texts.CARD_TITLE_UPDATE))
    assert "Talaba ma'lumotlarini yangiladi" in card_text
    assert "Guruh: DI-22" in card_text and "Turar joy: Kvartira" in card_text

    # pressing confirm again (double tap) neither duplicates rows nor notifications
    h.clear()
    await h.feed(callback_update(student, "reg:confirm:0"))
    assert h.answers()[-1].text == texts.STALE_BUTTON
    assert not h.cards(texts.CARD_TITLE_UPDATE) and not h.cards(texts.CARD_TITLE_NEW)
    assert len(await h.db.list_students()) == 1
    assert (await h.db.count_group_students(g1.id), await h.db.count_group_students(g2.id)) == (0, 1)


# ================================================================== (k) concurrent double-taps


async def test_concurrent_confirm_double_tap_saves_and_notifies_once(h: Harness) -> None:
    """Two ✅ Tasdiqlash taps handled concurrently (handle_as_tasks) must behave like one tap."""
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    group = await h.db.add_group(tutor.id, "DI-21")
    student = make_user(STUDENT)
    await register(h, student, tutor.id, group.id, confirm=False)
    h.clear()

    results = await h.feed_concurrently(callback_update(student, "reg:confirm:0"), callback_update(student, "reg:confirm:0"))
    assert not any(isinstance(r, BaseException) for r in results), results

    assert len(await h.db.list_students()) == 1
    assert await h.state_of(STUDENT) is None
    assert [m.text for m in h.messages(STUDENT)] == [texts.REG_SAVED]
    assert h.cards(texts.CARD_TITLE_NEW) == Counter({TUTOR: 1, SUPERADMIN: 1, BOTH: 1})
    assert not h.cards(texts.CARD_TITLE_UPDATE)
    # both taps answered: one silently, the other as a stale button
    assert sorted(str(a.text) for a in h.answers()) == sorted(["None", texts.STALE_BUTTON])
    assert_all_callbacks_answered(h)


async def test_concurrent_save_tutor_double_tap_adds_once(h: Harness) -> None:
    admin = make_user(SUPERADMIN)
    await h.feed(text_update(admin, "/add_tutor"), text_update(admin, "Karimov Aziz"), text_update(admin, str(TUTOR)))
    assert await h.state_of(SUPERADMIN) == AdminTutorAdd.confirm.state
    h.clear()

    results = await h.feed_concurrently(callback_update(admin, "adm:save:0"), callback_update(admin, "adm:save:0"))
    assert not any(isinstance(r, BaseException) for r in results), results

    assert [t.name for t in await h.db.list_tutors()] == ["Karimov Aziz"]
    assert await h.state_of(SUPERADMIN) is None and await h.data_of(SUPERADMIN) == {}
    shown = h.texts_to(SUPERADMIN)
    assert shown.count(texts.TUTOR_SAVED) == 1 and texts.TUTOR_TG_DUPLICATE not in shown
    assert_all_callbacks_answered(h)

    # the admin is not stuck: the next command works normally
    await h.feed(text_update(admin, "/tutors"))
    assert f"Karimov Aziz (id: {TUTOR})" in inline_buttons(h.last_message(SUPERADMIN).reply_markup).values()


async def test_concurrent_delete_double_tap_reports_second_as_gone(h: Harness) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    g1 = await h.db.add_group(tutor.id, "DI-21")
    tutor_user = make_user(TUTOR)
    await h.feed_concurrently(
        callback_update(tutor_user, f"tut:confirm_delete:{g1.id}:"), callback_update(tutor_user, f"tut:confirm_delete:{g1.id}:")
    )
    assert await h.db.get_group(g1.id) is None
    assert Counter(str(a.text) for a in h.answers()) == Counter({texts.GROUP_DELETED: 1, texts.GROUP_NOT_FOUND: 1})

    admin = make_user(SUPERADMIN)
    h.clear()
    await h.feed_concurrently(
        callback_update(admin, f"adm:confirm_delete:{tutor.id}"), callback_update(admin, f"adm:confirm_delete:{tutor.id}")
    )
    assert await h.db.get_tutor(tutor.id) is None
    assert Counter(str(a.text) for a in h.answers()) == Counter({texts.TUTOR_DELETED: 1, texts.TUTOR_NOT_FOUND: 1})
    assert_all_callbacks_answered(h)


# ================================================================== (l) commands typed mid-FSM


async def test_admin_commands_and_menu_buttons_preempt_admin_fsm(h: Harness) -> None:
    """A slash command or main-menu button typed at an admin prompt is never stored as data."""
    tutor = await h.db.add_tutor("Eski Ism", TUTOR)
    admin = make_user(SUPERADMIN)

    # add-tutor name prompt: every admin command pre-empts the step
    for command, expected in (
        ("/edit_tutor", texts.TUTOR_PICK_EDIT),
        ("/delete_tutor", texts.TUTOR_PICK_DELETE),
        ("/tutors", texts.TUTOR_LIST_TITLE.format(n=1)),
        ("/admin", texts.ADMIN_PANEL),
    ):
        await h.feed(text_update(admin, "/add_tutor"))
        assert await h.state_of(SUPERADMIN) == AdminTutorAdd.name.state
        await h.feed(text_update(admin, command))
        assert h.last_text(SUPERADMIN) == expected, command
        assert await h.state_of(SUPERADMIN) is None and await h.data_of(SUPERADMIN) == {}, command

    # ... and at the Telegram-ID prompt
    await h.feed(text_update(admin, "/add_tutor"), text_update(admin, "Yangi Tyutor"), text_update(admin, "/delete_tutor"))
    assert h.last_text(SUPERADMIN) == texts.TUTOR_PICK_DELETE
    assert await h.state_of(SUPERADMIN) is None
    assert [t.name for t in await h.db.list_tutors()] == ["Eski Ism"]

    # edit-name prompt: /delete_tutor must open the delete picker, not rename the tutor
    await h.feed(text_update(admin, "/edit_tutor"))
    kb = inline_buttons(h.last_message(SUPERADMIN).reply_markup)
    await h.feed(callback_update(admin, next(d for d, t in kb.items() if t.startswith("Eski Ism"))))
    field_kb = inline_buttons(h.last_shown(SUPERADMIN).reply_markup)
    name_btn = next(d for d, t in field_kb.items() if t == texts.BTN_EDIT_NAME)
    await h.feed(callback_update(admin, name_btn))
    assert await h.state_of(SUPERADMIN) == AdminTutorEdit.name.state
    await h.feed(text_update(admin, "/delete_tutor"))
    assert h.last_text(SUPERADMIN) == texts.TUTOR_PICK_DELETE
    unchanged = await h.db.get_tutor(tutor.id)
    assert unchanged is not None and unchanged.name == "Eski Ism"
    assert texts.TUTOR_UPDATED not in h.texts_to(SUPERADMIN)
    assert await h.state_of(SUPERADMIN) is None

    # a tutor-only command from a superadmin who is NOT a tutor is refused, not stored as the name
    await h.feed(text_update(admin, "/add_tutor"), text_update(admin, "/excel"))
    assert h.last_text(SUPERADMIN) == texts.TUTOR_ONLY
    assert await h.state_of(SUPERADMIN) == AdminTutorAdd.name.state and await h.data_of(SUPERADMIN) == {}
    await h.feed(text_update(admin, texts.BTN_TUTOR_PANEL))
    assert await h.data_of(SUPERADMIN) == {}
    # the register button does start registration from inside the admin FSM
    await h.feed(text_update(admin, texts.BTN_REGISTER), callback_update(admin, "sv:fill:basic"))
    assert h.last_text(SUPERADMIN) == texts.REG_CHOOSE_TUTOR
    assert await h.state_of(SUPERADMIN) == Registration.choose_tutor.state
    await h.feed(text_update(admin, "/cancel"))

    # a forwarded message is still accepted at the ID prompt even when its text is a command
    from aiogram.types import MessageOriginUser

    await h.feed(text_update(admin, "/add_tutor"), text_update(admin, "Forward Tyutor"))
    origin = MessageOriginUser(type="user", date=datetime.now(), sender_user=make_user(6006))
    await h.feed(Update(update_id=next(_update_ids), message=_message(admin, text="/tutor", forward_origin=origin)))
    assert await h.state_of(SUPERADMIN) == AdminTutorAdd.confirm.state
    assert (await h.data_of(SUPERADMIN))["telegram_id"] == 6006
    assert_all_callbacks_answered(h)


async def test_both_role_user_reaches_tutor_commands_from_admin_fsm(h: Harness) -> None:
    tutor = await h.db.add_tutor("Admin Tyutor", BOTH)
    await h.db.add_group(tutor.id, "AI-22")
    user = make_user(BOTH)

    for command, marker in (
        ("/tutor", "Tyutor panel"),
        ("/groups", "Guruhlarim"),
        ("/excel", texts.EXCEL_MENU),
        (texts.BTN_TUTOR_PANEL, "Tyutor panel"),
    ):
        for state in (AdminTutorAdd.name, AdminTutorAdd.confirm):
            await h.feed(text_update(user, "/add_tutor"))
            if state is AdminTutorAdd.confirm:
                await h.feed(text_update(user, "Kimdir Kimdirov"), text_update(user, "7007"))
            assert await h.state_of(BOTH) == state.state
            await h.feed(text_update(user, command))
            assert marker in h.last_text(BOTH), (command, state)
            assert await h.state_of(BOTH) is None and await h.data_of(BOTH) == {}, (command, state)
    assert [t.name for t in await h.db.list_tutors()] == ["Admin Tyutor"]
    # admin command from inside a tutor FSM step (the reverse direction)
    await h.feed(text_update(user, "/add_group"), text_update(user, "/add_tutor"))
    assert h.last_text(BOTH) == texts.ASK_TUTOR_NAME
    assert await h.state_of(BOTH) == AdminTutorAdd.name.state
    assert [g.name for g in await h.db.list_groups(tutor.id)] == ["AI-22"]


async def test_tutor_commands_preempt_tutor_fsm(h: Harness) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    g1 = await h.db.add_group(tutor.id, "G-1")
    tutor_user = make_user(TUTOR)

    for command, expected in (
        ("/excel", texts.EXCEL_MENU),
        ("/edit_group", texts.GROUP_PICK_EDIT),
        ("/delete_group", texts.GROUP_PICK_DELETE),
        ("/groups", texts.GROUP_LIST_TITLE.format(n=1)),
        # the panel also reminds a tutor who has not given their phone number yet
        ("/tutor", texts.TUTOR_PANEL.format(name="Karimov Aziz") + "\n\n" + texts.TUTOR_PHONE_MISSING),
    ):
        await h.feed(text_update(tutor_user, "/add_group"))
        assert await h.state_of(TUTOR) == TutorGroupAdd.name.state
        await h.feed(text_update(tutor_user, command))
        assert h.last_text(TUTOR) == expected, command
        assert await h.state_of(TUTOR) is None, command
    assert [g.name for g in await h.db.list_groups(tutor.id)] == ["G-1"]

    # an admin command from a plain tutor is refused instead of becoming a group name
    await h.feed(text_update(tutor_user, "/add_group"), text_update(tutor_user, "/admin"))
    assert h.last_text(TUTOR) == texts.ADMIN_ONLY
    assert await h.state_of(TUTOR) == TutorGroupAdd.name.state
    await h.feed(text_update(tutor_user, texts.BTN_ADMIN_PANEL))
    assert [g.name for g in await h.db.list_groups(tutor.id)] == ["G-1"]
    await h.feed(text_update(tutor_user, "/cancel"))

    # rename prompt: a command opens its own menu and the group keeps its name
    await h.feed(text_update(tutor_user, "/edit_group"), callback_update(tutor_user, f"tut:rename:{g1.id}:"))
    assert await h.state_of(TUTOR) == TutorGroupEdit.name.state
    await h.feed(text_update(tutor_user, "/delete_group"))
    assert h.last_text(TUTOR) == texts.GROUP_PICK_DELETE
    still = await h.db.get_group(g1.id)
    assert still is not None and still.name == "G-1"
    assert await h.state_of(TUTOR) is None
    assert texts.GROUP_RENAMED not in h.texts_to(TUTOR)


# ================================================================== extra: robustness


async def test_group_deleted_before_confirm_restarts(h: Harness) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    group = await h.db.add_group(tutor.id, "DI-21")
    student = make_user(STUDENT)
    await register(h, student, tutor.id, group.id, confirm=False)
    await h.db.delete_group(group.id)
    h.clear()
    await h.feed(callback_update(student, "reg:confirm:0"))
    assert h.answers()[0].text == texts.REG_GROUP_NOT_FOUND
    assert h.last_text(STUDENT) == texts.REG_CHOOSE_TUTOR
    assert await h.state_of(STUDENT) == Registration.choose_tutor.state
    assert await h.db.get_student_by_telegram_id(STUDENT) is None
    assert not h.cards(texts.CARD_TITLE_NEW)


async def test_tutor_without_groups_and_back_navigation(h: Harness) -> None:
    empty = await h.db.add_tutor("Bo'sh Tyutor", OTHER_TUTOR)
    full = await h.db.add_tutor("Karimov Aziz", TUTOR)
    group = await h.db.add_group(full.id, "DI-21")
    student = make_user(STUDENT)

    await h.feed(
        text_update(student, "/start"),
        callback_update(student, "sv:fill:basic"),
        callback_update(student, f"reg:tutor:{empty.id}"),
    )
    assert h.last_text(STUDENT) == texts.REG_TUTOR_NO_GROUPS
    kb = inline_buttons(h.last_shown(STUDENT).reply_markup)
    assert "reg:back:0" in kb and not any(d.startswith("reg:group:") for d in kb)
    await h.feed(callback_update(student, "reg:back:0"))
    assert h.last_text(STUDENT) == texts.REG_CHOOSE_TUTOR
    assert await h.state_of(STUDENT) == Registration.choose_tutor.state

    # a group id of another tutor cannot be smuggled in via callback data
    await h.feed(callback_update(student, f"reg:tutor:{empty.id}"))
    await h.feed(callback_update(student, f"reg:group:{group.id}"))
    assert h.answers()[-1].text == texts.REG_GROUP_NOT_FOUND
    assert "group_id" not in await h.data_of(STUDENT)

    # a deleted tutor id
    await h.feed(callback_update(student, "reg:tutor:9999"))
    assert h.answers()[-1].text == texts.REG_TUTOR_NOT_FOUND
    assert_all_callbacks_answered(h)


async def test_notification_failure_does_not_break_confirmation(h: Harness, monkeypatch: pytest.MonkeyPatch) -> None:
    from aiogram.exceptions import TelegramForbiddenError

    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    group = await h.db.add_group(tutor.id, "DI-21")
    student = make_user(STUDENT)
    original = h.session.make_request

    async def failing(bot: Bot, method: TelegramMethod[Any], timeout: int | None = None) -> Any:
        if type(method).__name__ == "SendMessage" and int(method.chat_id) == TUTOR:
            h.session.requests.append(("SendMessage", method))
            raise TelegramForbiddenError(method=method, message="Forbidden: bot was blocked by the user")
        return await original(bot, method, timeout)

    monkeypatch.setattr(h.session, "make_request", failing)
    await register(h, student, tutor.id, group.id)

    assert (await h.db.get_student_by_telegram_id(STUDENT)) is not None
    assert texts.REG_SAVED in h.texts_to(STUDENT)
    assert h.cards(texts.CARD_TITLE_NEW) == Counter({TUTOR: 1, SUPERADMIN: 1, BOTH: 1})
    assert await h.state_of(STUDENT) is None


# ================================================================== extra: commands, cascades, edits


async def test_setup_bot_commands_scopes(h: Harness) -> None:
    from bot.commands import setup_bot_commands

    await h.db.add_tutor("Admin Tyutor", BOTH)
    await h.db.add_tutor("Karimov Aziz", TUTOR)
    await setup_bot_commands(h.bot, h.db, h.settings)

    by_scope: dict[Any, set[str]] = {}
    for m in h.of("SetMyCommands"):
        key = getattr(m.scope, "chat_id", None) if m.scope is not None else "default"
        if key is None:
            key = type(m.scope).__name__
        by_scope[key] = {c.command for c in m.commands}
    assert by_scope.get("BotCommandScopeDefault") == {"start", "mydata", "help", "cancel"}
    assert by_scope[SUPERADMIN] == {
        "start",
        "mydata",
        "help",
        "cancel",
        "admin",
        "tutors",
        "users",
        "test_users",
        "add_tutor",
        "edit_tutor",
        "delete_tutor",
        "broadcast",
    }
    assert {"tutor", "groups", "add_group", "edit_group", "delete_group", "excel"} <= by_scope[TUTOR]
    assert "admin" not in by_scope[TUTOR]
    assert {"admin", "tutor", "add_tutor", "add_group"} <= by_scope[BOTH]  # additive for BOTH
    assert STUDENT not in by_scope and OTHER_TUTOR not in by_scope


async def test_admin_delete_tutor_cascades_and_revokes_access(h: Harness) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    group = await h.db.add_group(tutor.id, "DI-21")
    await register(h, make_user(STUDENT), tutor.id, group.id)
    tutor_user = make_user(TUTOR)
    await h.feed(text_update(tutor_user, "/tutor"))
    assert "Tyutor panel" in h.last_text(TUTOR)
    h.clear()

    admin = make_user(SUPERADMIN)
    await h.feed(text_update(admin, "/delete_tutor"))
    kb = inline_buttons(h.last_message(SUPERADMIN).reply_markup)
    pick = next(d for d, t in kb.items() if t == f"Karimov Aziz (id: {TUTOR})")
    await h.feed(callback_update(admin, pick))
    assert "1 ta guruh va 1 ta talaba ham o'chiriladi" in h.last_text(SUPERADMIN)
    confirm_kb = inline_buttons(h.last_shown(SUPERADMIN).reply_markup)
    yes = next(d for d, t in confirm_kb.items() if t == texts.BTN_YES_DELETE)
    assert texts.BTN_NO in confirm_kb.values()
    assert await h.db.get_tutor(tutor.id) is not None  # nothing deleted before confirmation

    await h.feed(callback_update(admin, yes))
    assert await h.db.get_tutor(tutor.id) is None
    assert await h.db.get_group(group.id) is None
    assert await h.db.get_student_by_telegram_id(STUDENT) is None
    assert texts.TUTOR_DELETED in h.last_text(SUPERADMIN)
    # the removed tutor's menu is reset and they lose access on the next action (live check)
    assert TUTOR in [getattr(m.scope, "chat_id", None) for m in h.of("DeleteMyCommands")]
    await h.feed(text_update(tutor_user, "/tutor"))
    assert h.last_text(TUTOR) == texts.TUTOR_ONLY
    assert_all_callbacks_answered(h)


async def test_tutor_delete_group_cascades(h: Harness) -> None:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    g1 = await h.db.add_group(tutor.id, "DI-21")
    g2 = await h.db.add_group(tutor.id, "DI-22")
    await register(h, make_user(STUDENT), tutor.id, g1.id)
    await register(h, make_user(4005), tutor.id, g2.id, RegInput(phone="+998907777777"))
    h.clear()
    tutor_user = make_user(TUTOR)

    await h.feed(text_update(tutor_user, "/delete_group"))
    kb = inline_buttons(h.last_message(TUTOR).reply_markup)
    pick = next(d for d, t in kb.items() if t.startswith("DI-21"))
    assert kb[pick] == "DI-21 — 📋 1 ta · 🗂 0 ta"
    await h.feed(callback_update(tutor_user, pick))
    assert "📋 1 ta asosiy va 🗂 0 ta to'liq anketa" in h.last_text(TUTOR)
    confirm_kb = inline_buttons(h.last_shown(TUTOR).reply_markup)
    yes = next(d for d, t in confirm_kb.items() if t == texts.BTN_YES_DELETE)
    await h.feed(callback_update(tutor_user, yes))
    assert await h.db.get_group(g1.id) is None
    assert await h.db.get_student_by_telegram_id(STUDENT) is None
    assert await h.db.get_student_by_telegram_id(4005) is not None  # other group untouched
    assert texts.GROUP_DELETED in h.last_text(TUTOR)
    assert_all_callbacks_answered(h)


async def test_admin_edit_tutor_name_and_forwarded_telegram_id(h: Harness) -> None:
    from aiogram.types import MessageOriginHiddenUser, MessageOriginUser

    tutor = await h.db.add_tutor("Eski Ism", TUTOR)
    await h.db.add_tutor("Boshqa Tyutor", OTHER_TUTOR)
    admin = make_user(SUPERADMIN)

    await h.feed(text_update(admin, "/edit_tutor"))
    kb = inline_buttons(h.last_message(SUPERADMIN).reply_markup)
    pick = next(d for d, t in kb.items() if t.startswith("Eski Ism"))
    await h.feed(callback_update(admin, pick))
    field_kb = inline_buttons(h.last_shown(SUPERADMIN).reply_markup)
    name_btn = next(d for d, t in field_kb.items() if t == texts.BTN_EDIT_NAME)
    tg_btn = next(d for d, t in field_kb.items() if t == texts.BTN_EDIT_TG)

    await h.feed(callback_update(admin, name_btn), text_update(admin, "Yangi Ism"))
    updated = await h.db.get_tutor(tutor.id)
    assert updated is not None and updated.name == "Yangi Ism"
    assert texts.TUTOR_UPDATED in h.texts_to(SUPERADMIN)
    assert await h.state_of(SUPERADMIN) is None

    # telegram id: duplicate refused, hidden forward explained, visible forward accepted
    await h.feed(callback_update(admin, tg_btn), text_update(admin, str(OTHER_TUTOR)))
    assert h.last_text(SUPERADMIN) == texts.TUTOR_TG_DUPLICATE
    unchanged = await h.db.get_tutor(tutor.id)
    assert unchanged is not None and unchanged.telegram_id == TUTOR
    for bad in ("99999999999999999999", "²"):
        await h.feed(text_update(admin, bad))
        assert h.last_text(SUPERADMIN) == texts.TUTOR_TG_INVALID, bad
        assert await h.state_of(SUPERADMIN) == AdminTutorEdit.telegram_id.state

    hidden_origin = MessageOriginHiddenUser(type="hidden_user", date=datetime.now(), sender_user_name="X")
    await h.feed(Update(update_id=next(_update_ids), message=_message(admin, text="hi", forward_origin=hidden_origin)))
    assert h.last_text(SUPERADMIN) == texts.TUTOR_TG_HIDDEN

    origin = MessageOriginUser(type="user", date=datetime.now(), sender_user=make_user(6006))
    await h.feed(Update(update_id=next(_update_ids), message=_message(admin, text="hi", forward_origin=origin)))
    updated = await h.db.get_tutor(tutor.id)
    assert updated is not None and updated.telegram_id == 6006
    assert texts.TUTOR_UPDATED in h.texts_to(SUPERADMIN)
    # command menus of the old and the new chat are refreshed
    touched = {getattr(m.scope, "chat_id", None) for m in h.of("SetMyCommands") + h.of("DeleteMyCommands")}
    assert {TUTOR, 6006} <= touched
    # old id lost access, new id gained it
    await h.feed(text_update(make_user(TUTOR), "/tutor"))
    assert h.last_text(TUTOR) == texts.TUTOR_ONLY
    await h.feed(text_update(make_user(6006), "/tutor"))
    assert "Tyutor panel" in h.last_text(6006)


async def test_role_user_finishing_registration_gets_main_menu(h: Harness) -> None:
    tutor = await h.db.add_tutor("Admin Tyutor", BOTH)
    group = await h.db.add_group(tutor.id, "AI-22")
    user = make_user(BOTH)
    await register(h, user, tutor.id, group.id, start_text=texts.BTN_REGISTER)
    saved = next(m for m in h.messages(BOTH) if m.text == texts.REG_SAVED)
    assert set(reply_button_texts(saved.reply_markup)) == {
        texts.BTN_ADMIN_PANEL,
        texts.BTN_TUTOR_PANEL,
        texts.BTN_REGISTER,
    }
    assert h.cards(texts.CARD_TITLE_NEW) == Counter({BOTH: 1, SUPERADMIN: 1})
    assert (await h.db.get_student_by_telegram_id(BOTH)) is not None


async def test_help_is_role_specific(h: Harness) -> None:
    await h.db.add_tutor("Karimov Aziz", TUTOR)
    await h.feed(text_update(make_user(STUDENT), "/help"))
    student_help = h.last_text(STUDENT)
    assert "/start" in student_help and "/add_tutor" not in student_help and "/add_group" not in student_help
    await h.feed(text_update(make_user(TUTOR), "/help"))
    assert "/add_group" in h.last_text(TUTOR) and "/add_tutor" not in h.last_text(TUTOR)
    await h.feed(text_update(make_user(SUPERADMIN), "/help"))
    assert "/add_tutor" in h.last_text(SUPERADMIN) and "/add_group" not in h.last_text(SUPERADMIN)
