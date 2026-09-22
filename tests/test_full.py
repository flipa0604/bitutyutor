"""The full questionnaire: filling it in, the checks that make faking it hard, editing and exports.

Runs through the real Dispatcher with the flow harness from ``tests/test_flows.py``.
User ids: SUPERADMIN=1001, TUTOR=2002, BOTH=3003 (superadmin AND tutor), STUDENT=4004, OTHER_TUTOR=5005.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from bot import texts
from bot.excel import FULL_HEADERS
from bot.states import FullSurvey
from tests.test_flows import (
    OTHER_TUTOR,
    STUDENT,
    SUPERADMIN,
    TUTOR,
    Harness,
    assert_all_callbacks_answered,
    callback_update,
    contact_update,
    h,
    inline_buttons,
    make_user,
    reply_button_texts,
    text_update,
)

__all__ = ["h"]  # the flow harness fixture is re-exported so pytest finds it in this module too

SECOND_STUDENT = 4005

PHONE = "+998901234567"
BIRTH = "05.03.2004"
PASSPORT = "AA1234567"
PINFL = "50503042340012"  # 5 = male born in the 2000s, then 050304 = the birth date above

ANSWERS: dict[str, Any] = {
    "full_name": "Aliyev Vali G'aniyevich",
    "direction": "Dasturiy injiniring",
    "course": "3-kurs",
    "birth_date": BIRTH,
    "passport": PASSPORT,
    "pinfl": PINFL,
    "citizenship": texts.BTN_CITIZEN_UZ,
    "region": "Toshkent shahri",
    "district": "Chilonzor tumani",
    "mfy": "Navbahor MFY",
    "mfy_contact": "+998901112233",
    "street": "Navoiy ko'chasi, 12-uy",
    "employed": texts.BTN_EMPLOYED_NO,
    "married": texts.BTN_MARRIED_NO,
    "father_name": "Aliyev G'ani Karimovich",
    "father_phone": "+998901111111",
    "father_work": "Maktab",
    "mother_name": "Aliyeva Zulfiya Anvarovna",
    "mother_phone": "+998902222222",
    "mother_work": "Uy bekasi",
}
"""A complete set of answers for the shortest path (no job, no spouse, no social status)."""


async def open_full(h: Harness, user: Any, tutor_id: int, group_id: int) -> None:
    """/start → 🗂 To'liq anketa → tutor → group, i.e. everything before the first question."""
    await h.feed(
        text_update(user, "/start"),
        callback_update(user, "sv:fill:full"),
        callback_update(user, f"reg:tutor:{tutor_id}"),
        callback_update(user, f"reg:group:{group_id}"),
    )


async def fill_full(
    h: Harness,
    user: Any,
    tutor_id: int,
    group_id: int,
    *,
    answers: dict[str, Any] | None = None,
    social: list[str] | None = None,
    confirm: bool = True,
) -> None:
    """Answer every question of the full survey in order, then confirm the preview."""
    values = {**ANSWERS, **(answers or {})}
    await open_full(h, user, tutor_id, group_id)
    await h.feed(contact_update(user, values.get("phone", PHONE), user.id))
    for field in ("full_name", "direction", "course", "birth_date", "passport", "pinfl", "citizenship"):
        await h.feed(text_update(user, values[field]))
    if values["citizenship"] == texts.BTN_CITIZEN_OTHER:
        await h.feed(text_update(user, values["citizenship_other"]))
    for field in ("region", "district", "mfy", "mfy_contact", "street", "employed"):
        await h.feed(text_update(user, values[field]))
    if values["employed"] == texts.BTN_EMPLOYED_YES:
        for field in ("work_place", "work_position", "work_address", "work_phone"):
            await h.feed(text_update(user, values[field]))
    await h.feed(text_update(user, values["married"]))
    if values["married"] == texts.BTN_MARRIED_YES:
        for field in ("spouse_name", "spouse_work", "spouse_phone"):
            await h.feed(text_update(user, values[field]))
    for code in social or []:
        await h.feed(callback_update(user, f"fl:soc:{code}"))
    await h.feed(callback_update(user, "fl:soc_ok:"))
    for field in ("father_name", "father_phone", "father_work", "mother_name", "mother_phone", "mother_work"):
        await h.feed(text_update(user, values[field]))
    assert texts.FULL_PREVIEW_TITLE in h.last_text(user.id)
    if confirm:
        await h.feed(callback_update(user, "fl:confirm:"))


async def seed(h: Harness) -> tuple[Any, Any]:
    tutor = await h.db.add_tutor("Karimov Aziz", TUTOR)
    group = await h.db.add_group(tutor.id, "DI-21")
    return tutor, group


# ================================================================== the picker


async def test_start_offers_both_questionnaires_independently(h: Harness) -> None:
    tutor, group = await seed(h)
    student = make_user(STUDENT, username="vali")

    await h.feed(text_update(student, "/start"))
    assert inline_buttons(h.last_message(STUDENT).reply_markup) == {
        "sv:fill:basic": f"🕗 {texts.BTN_SURVEY_BASIC}",
        "sv:fill:full": f"🕗 {texts.BTN_SURVEY_FULL}",
    }

    await fill_full(h, student, tutor.id, group.id)
    assert await h.db.get_full_profile_by_telegram_id(STUDENT) is not None
    assert await h.db.get_student_by_telegram_id(STUDENT) is None  # the basic one stays untouched

    # the picker now marks the full one as done and still offers the basic one
    await h.feed(text_update(student, "/start"))
    buttons = inline_buttons(h.last_message(STUDENT).reply_markup)
    assert buttons == {
        "sv:fill:basic": f"🕗 {texts.BTN_SURVEY_BASIC}",
        "sv:open:full": f"✅ {texts.BTN_SURVEY_FULL}",
    }
    assert "to'ldirilgan" in h.last_text(STUDENT) and "to'ldirilmagan" in h.last_text(STUDENT)

    # opening the finished one shows the card with its edit button
    await h.feed(callback_update(student, "sv:open:full"))
    card = h.last_message(STUDENT)
    assert card.text.startswith(texts.FULL_HOME_TITLE)
    assert inline_buttons(card.reply_markup) == {"fl:open:": texts.BTN_EDIT_FULL_DATA}
    assert_all_callbacks_answered(h)


# ================================================================== filling it in


async def test_full_survey_saves_every_answer_and_notifies(h: Harness) -> None:
    tutor, group = await seed(h)
    student = make_user(STUDENT, username="vali")

    await open_full(h, student, tutor.id, group.id)
    prompt = h.last_message(STUDENT)
    assert texts.FULL_ASK_PHONE in prompt.text and "1/22" in prompt.text
    assert [b.text for row in prompt.reply_markup.keyboard for b in row] == [
        texts.BTN_SEND_CONTACT,
        texts.BTN_BACK,
        texts.BTN_CANCEL,
    ]
    assert await h.state_of(STUDENT) == FullSurvey.phone.state

    await fill_full(h, student, tutor.id, group.id, social=["yetim", "kam_taminlangan"])

    profile = await h.db.get_full_profile_by_telegram_id(STUDENT)
    assert profile is not None
    assert (profile.full_name, profile.phone, profile.direction) == (ANSWERS["full_name"], PHONE, "Dasturiy injiniring")
    assert (profile.course, profile.birth_date, profile.passport, profile.pinfl) == (3, BIRTH, PASSPORT, PINFL)
    assert (profile.citizenship, profile.region, profile.district) == ("O'zbekiston", "Toshkent shahri", "Chilonzor tumani")
    assert (profile.mfy, profile.mfy_contact, profile.street) == ("Navbahor MFY", "+998901112233", "Navoiy ko'chasi, 12-uy")
    assert profile.employed is False and profile.work_place == ""
    assert profile.married is False and profile.spouse_name == ""
    assert profile.social_codes == ["yetim", "kam_taminlangan"]
    assert (profile.father_name, profile.father_phone, profile.father_work) == (
        ANSWERS["father_name"], "+998901111111", "Maktab",
    )
    assert (profile.mother_name, profile.mother_phone, profile.mother_work) == (
        ANSWERS["mother_name"], "+998902222222", "Uy bekasi",
    )
    assert profile.tutor_id == tutor.id and profile.group_id == group.id and profile.username == "vali"

    assert h.last_text(STUDENT) == texts.FULL_SAVED
    assert await h.state_of(STUDENT) is None
    cards = [m for m in h.of("SendMessage") if str(m.text or "").startswith(texts.FULL_CARD_TITLE_NEW)]
    assert sorted(m.chat_id for m in cards) == [SUPERADMIN, TUTOR, 3003]  # tutor + both superadmins
    assert "Yetim, Kam ta'minlangan" in str(cards[0].text)
    assert_all_callbacks_answered(h)


async def test_work_and_family_blocks_are_asked_only_when_they_apply(h: Harness) -> None:
    tutor, group = await seed(h)
    student = make_user(STUDENT)

    await fill_full(
        h,
        student,
        tutor.id,
        group.id,
        answers={
            "employed": texts.BTN_EMPLOYED_YES,
            "work_place": "IT Park",
            "work_position": "Dasturchi",
            "work_address": "Toshkent, Mirzo Ulug'bek",
            "work_phone": "+998712001122",
            "married": texts.BTN_MARRIED_YES,
            "spouse_name": "Aliyeva Nilufar Bekzodovna",
            "spouse_work": "Shifokor",
            "spouse_phone": "+998903334455",
            "citizenship": texts.BTN_CITIZEN_OTHER,
            "citizenship_other": "Qozog'iston",
        },
    )
    profile = await h.db.get_full_profile_by_telegram_id(STUDENT)
    assert profile is not None
    assert profile.employed and (profile.work_place, profile.work_position) == ("IT Park", "Dasturchi")
    assert (profile.work_address, profile.work_phone) == ("Toshkent, Mirzo Ulug'bek", "+998712001122")
    assert profile.married and profile.spouse_name == "Aliyeva Nilufar Bekzodovna"
    assert (profile.spouse_work, profile.spouse_phone) == ("Shifokor", "+998903334455")
    assert profile.citizenship == "Qozog'iston"
    assert profile.social_codes == []
    assert_all_callbacks_answered(h)


async def test_orphan_may_skip_the_parent_fields(h: Harness) -> None:
    tutor, group = await seed(h)
    student = make_user(STUDENT)
    skip = texts.BTN_FULL_SKIP_PARENT
    await fill_full(
        h,
        student,
        tutor.id,
        group.id,
        social=["chin_yetim"],
        answers={f"{parent}_{field}": skip for parent in ("father", "mother") for field in ("name", "phone", "work")},
    )
    profile = await h.db.get_full_profile_by_telegram_id(STUDENT)
    assert profile is not None
    assert profile.father_name == texts.NO_DATA_VALUE and profile.mother_phone == texts.NO_DATA_VALUE
    assert profile.social_codes == ["chin_yetim"]


# ================================================================== the checks


async def test_the_phone_must_come_from_the_contact_button(h: Harness) -> None:
    tutor, group = await seed(h)
    student = make_user(STUDENT)
    await open_full(h, student, tutor.id, group.id)

    await h.feed(text_update(student, "+998901234567"))  # typed: proves nothing
    assert h.last_text(STUDENT) == texts.FULL_PHONE_BUTTON_ONLY
    assert await h.state_of(STUDENT) == FullSurvey.phone.state

    await h.feed(contact_update(student, "+998907654321", SECOND_STUDENT))  # somebody else's card
    assert h.last_text(STUDENT) == texts.REG_CONTACT_NOT_OWN
    assert await h.state_of(STUDENT) == FullSurvey.phone.state

    await h.feed(contact_update(student, "998901234567", STUDENT))
    assert texts.FULL_ASK_FULL_NAME in h.last_text(STUDENT)
    assert (await h.data_of(STUDENT))["answers"]["phone"] == "+998901234567"


async def test_passport_pinfl_and_birth_date_are_checked_against_each_other(h: Harness) -> None:
    tutor, group = await seed(h)
    student = make_user(STUDENT)
    await open_full(h, student, tutor.id, group.id)
    await h.feed(contact_update(student, PHONE, STUDENT))
    for field in ("full_name", "direction", "course"):
        await h.feed(text_update(student, ANSWERS[field]))

    await h.feed(text_update(student, "31.02.2004"))  # a day that does not exist
    assert h.last_text(STUDENT).startswith("❗️ Sana noto'g'ri")
    await h.feed(text_update(student, "05.03.2019"))  # a five-year-old student
    assert h.last_text(STUDENT).startswith("❗️ Sana noto'g'ri")
    await h.feed(text_update(student, BIRTH))
    assert texts.FULL_ASK_PASSPORT in h.last_text(STUDENT)

    await h.feed(text_update(student, "A1234567"))  # one letter short
    assert h.last_text(STUDENT) == texts.FULL_PASSPORT_INVALID
    await h.feed(text_update(student, "aa 1234567"))  # spacing and case are forgiven
    assert texts.FULL_ASK_PINFL in h.last_text(STUDENT)
    assert (await h.data_of(STUDENT))["answers"]["passport"] == PASSPORT

    await h.feed(text_update(student, "12345678901234"))  # random 14 digits
    assert h.last_text(STUDENT).startswith("❗️ JShShR noto'g'ri")
    await h.feed(text_update(student, "51103042340012"))  # right shape, wrong birth date inside
    assert h.last_text(STUDENT).startswith("❗️ JShShR noto'g'ri")
    await h.feed(text_update(student, "30503042340012"))  # born in the 1900s? not with that date
    assert h.last_text(STUDENT).startswith("❗️ JShShR noto'g'ri")
    await h.feed(text_update(student, PINFL))
    assert texts.FULL_ASK_CITIZENSHIP in h.last_text(STUDENT)


async def test_one_passport_cannot_be_used_twice(h: Harness) -> None:
    tutor, group = await seed(h)
    first = make_user(STUDENT)
    await fill_full(h, first, tutor.id, group.id)
    assert await h.db.get_full_profile_by_telegram_id(STUDENT) is not None

    second = make_user(SECOND_STUDENT)
    await open_full(h, second, tutor.id, group.id)
    await h.feed(contact_update(second, "+998905554433", SECOND_STUDENT))
    for field in ("full_name", "direction", "course", "birth_date", "passport"):
        await h.feed(text_update(second, ANSWERS[field]))
    await h.feed(text_update(second, PINFL))
    assert h.last_text(SECOND_STUDENT) == texts.FULL_IDENTITY_TAKEN
    assert await h.state_of(SECOND_STUDENT) == FullSurvey.pinfl.state
    assert await h.db.get_full_profile_by_telegram_id(SECOND_STUDENT) is None


async def test_choices_are_taken_from_the_buttons_only(h: Harness) -> None:
    tutor, group = await seed(h)
    student = make_user(STUDENT)
    await open_full(h, student, tutor.id, group.id)
    await h.feed(contact_update(student, PHONE, STUDENT))
    await h.feed(text_update(student, "Aliyev"))  # one word is not a full name
    assert h.last_text(STUDENT) == texts.REG_NAME_INVALID
    await h.feed(text_update(student, ANSWERS["full_name"]), text_update(student, ANSWERS["direction"]))

    course_prompt = h.last_message(STUDENT)
    assert reply_button_texts(course_prompt.reply_markup) == [
        "1-kurs", "2-kurs", "3-kurs", "4-kurs", "5-kurs", "6-kurs", texts.BTN_BACK, texts.BTN_CANCEL,
    ]
    await h.feed(text_update(student, "9-kurs"))
    assert h.last_text(STUDENT) == texts.FULL_COURSE_INVALID
    await h.feed(text_update(student, "3-kurs"))
    assert texts.FULL_ASK_BIRTH in h.last_text(STUDENT)

    for field in ("birth_date", "passport", "pinfl", "citizenship"):
        await h.feed(text_update(student, ANSWERS[field]))
    region_prompt = h.last_message(STUDENT)
    assert "Andijon" in reply_button_texts(region_prompt.reply_markup)
    await h.feed(text_update(student, "Marsdagi viloyat"))
    assert h.last_text(STUDENT) == texts.FULL_REGION_INVALID
    await h.feed(text_update(student, "toshkent shahri"))  # the button's text, case-insensitively
    assert (await h.data_of(STUDENT))["answers"]["region"] == "Toshkent shahri"


async def test_going_back_re_asks_the_previous_question(h: Harness) -> None:
    tutor, group = await seed(h)
    student = make_user(STUDENT)
    await open_full(h, student, tutor.id, group.id)
    await h.feed(contact_update(student, PHONE, STUDENT))
    await h.feed(text_update(student, ANSWERS["full_name"]), text_update(student, ANSWERS["direction"]))
    assert await h.state_of(STUDENT) == FullSurvey.course.state

    await h.feed(text_update(student, texts.BTN_BACK))
    assert texts.FULL_ASK_DIRECTION in h.last_text(STUDENT)
    assert await h.state_of(STUDENT) == FullSurvey.direction.state
    await h.feed(text_update(student, texts.BTN_BACK))
    assert texts.FULL_ASK_FULL_NAME in h.last_text(STUDENT)
    await h.feed(text_update(student, "Yangi Ism Otasi"))  # answer it differently this time
    assert await h.state_of(STUDENT) == FullSurvey.direction.state
    assert (await h.data_of(STUDENT))["answers"]["full_name"] == "Yangi Ism Otasi"

    # the very first question has nothing behind it
    await h.feed(text_update(student, texts.BTN_BACK), text_update(student, texts.BTN_BACK))
    assert await h.state_of(STUDENT) == FullSurvey.phone.state
    await h.feed(text_update(student, texts.BTN_BACK))
    assert h.last_text(STUDENT) == texts.FULL_NOTHING_TO_GO_BACK


async def test_cancel_and_restart(h: Harness) -> None:
    tutor, group = await seed(h)
    student = make_user(STUDENT)
    await open_full(h, student, tutor.id, group.id)
    await h.feed(contact_update(student, PHONE, STUDENT), text_update(student, texts.BTN_CANCEL))
    assert h.last_text(STUDENT) == texts.CANCELLED_STUDENT
    assert await h.state_of(STUDENT) is None
    assert await h.db.get_full_profile_by_telegram_id(STUDENT) is None

    await fill_full(h, student, tutor.id, group.id, confirm=False)
    await h.feed(callback_update(student, "fl:restart:"))
    assert texts.FULL_RESTARTED in h.texts_to(STUDENT)
    assert await h.state_of(STUDENT) == FullSurvey.choose_tutor.state
    assert await h.db.get_full_profile_by_telegram_id(STUDENT) is None

    await h.feed(callback_update(student, f"reg:tutor:{tutor.id}"), callback_update(student, "fl:cancel:"))
    assert h.last_text(STUDENT) == texts.CANCELLED_STUDENT
    assert await h.state_of(STUDENT) is None
    assert_all_callbacks_answered(h)


# ================================================================== editing


async def test_student_edits_single_fields(h: Harness) -> None:
    tutor, group = await seed(h)
    student = make_user(STUDENT)
    await fill_full(h, student, tutor.id, group.id)
    h.clear()

    await h.feed(text_update(student, "/mydata"), callback_update(student, "sv:open:full"))
    await h.feed(callback_update(student, "fl:open:"))
    menu = h.last_message(STUDENT)
    assert texts.FULL_EDIT_MENU in menu.text
    buttons = inline_buttons(menu.reply_markup)
    assert {"fl:field:full_name", "fl:field:street", "fl:field:employed", "fl:done:"} <= set(buttons)

    await h.feed(callback_update(student, "fl:field:street"), text_update(student, "Amir Temur ko'chasi, 1-uy"))
    profile = await h.db.get_full_profile_by_telegram_id(STUDENT)
    assert profile is not None and profile.street == "Amir Temur ko'chasi, 1-uy"
    assert profile.edited_at and texts.EDIT_SAVED in h.texts_to(STUDENT)

    # the work block asks its sub-questions when the answer turns to "yes"
    await h.feed(callback_update(student, "fl:field:employed"), text_update(student, texts.BTN_EMPLOYED_YES))
    for value in ("IT Park", "Dasturchi", "Toshkent", "+998712001122"):
        await h.feed(text_update(student, value))
    profile = await h.db.get_full_profile_by_telegram_id(STUDENT)
    assert profile is not None and profile.employed and profile.work_place == "IT Park"

    # ... and clears them again when it turns back to "no"
    await h.feed(callback_update(student, "fl:field:employed"), text_update(student, texts.BTN_EMPLOYED_NO))
    profile = await h.db.get_full_profile_by_telegram_id(STUDENT)
    assert profile is not None and not profile.employed and profile.work_place == ""

    await h.feed(callback_update(student, "fl:done:"))
    assert texts.EDIT_DONE in h.texts_to(STUDENT)
    assert await h.state_of(STUDENT) is None
    updates = [m for m in h.of("SendMessage") if str(m.text or "").startswith(texts.FULL_CARD_TITLE_UPDATE)]
    assert sorted(m.chat_id for m in updates) == [SUPERADMIN, TUTOR, 3003]
    assert_all_callbacks_answered(h)


async def test_changing_the_birth_date_re_asks_the_pinfl(h: Harness) -> None:
    tutor, group = await seed(h)
    student = make_user(STUDENT)
    await fill_full(h, student, tutor.id, group.id)
    await h.feed(text_update(student, "/mydata"), callback_update(student, "sv:open:full"), callback_update(student, "fl:open:"))

    await h.feed(callback_update(student, "fl:field:birth_date"), text_update(student, "07.07.2003"))
    assert texts.FULL_ASK_PINFL in h.last_text(STUDENT)
    await h.feed(text_update(student, PINFL))  # the old one no longer matches the new date
    assert h.last_text(STUDENT).startswith("❗️ JShShR noto'g'ri")
    await h.feed(text_update(student, "50707032340012"))
    profile = await h.db.get_full_profile_by_telegram_id(STUDENT)
    assert profile is not None and profile.birth_date == "07.07.2003" and profile.pinfl == "50707032340012"


async def test_student_moves_to_another_tutor(h: Harness) -> None:
    tutor, group = await seed(h)
    other = await h.db.add_tutor("Boshqa Tyutor", OTHER_TUTOR)
    other_group = await h.db.add_group(other.id, "IQ-11")
    student = make_user(STUDENT)
    await fill_full(h, student, tutor.id, group.id)
    await h.feed(text_update(student, "/mydata"), callback_update(student, "sv:open:full"), callback_update(student, "fl:open:"))

    await h.feed(
        callback_update(student, "fl:field:tg"),
        callback_update(student, f"reg:tutor:{other.id}"),
        callback_update(student, f"reg:group:{other_group.id}"),
    )
    profile = await h.db.get_full_profile_by_telegram_id(STUDENT)
    assert profile is not None and profile.tutor_id == other.id and profile.group_id == other_group.id
    assert_all_callbacks_answered(h)


# ================================================================== exports


async def test_tutor_exports_the_full_survey(h: Harness) -> None:
    tutor, group = await seed(h)
    student = make_user(STUDENT, username="vali")
    await fill_full(h, student, tutor.id, group.id, social=["nogiron"])
    tutor_user = make_user(TUTOR)
    await h.db.update_tutor_phone(tutor.id, "+998909998877")
    h.clear()

    await h.feed(text_update(tutor_user, "/excel"))
    menu = inline_buttons(h.last_message(TUTOR).reply_markup)
    assert menu["tut:exf_pick:0:"] == texts.BTN_EXCEL_FULL_GROUP
    assert menu["tut:exf_all:0:"] == texts.BTN_EXCEL_FULL_ALL

    await h.feed(callback_update(tutor_user, "tut:exf_all:0:"))
    docs = h.documents(TUTOR)
    assert len(docs) == 1 and docs[0].document.filename.endswith(".xlsx")
    wb = load_workbook(BytesIO(docs[0].document.data))
    assert wb.sheetnames == ["Barchasi", "DI-21"]
    rows = [list(r) for r in wb["Barchasi"].iter_rows(values_only=True)]
    assert rows[0] == list(FULL_HEADERS)
    row = rows[1]
    assert row[1] == "Karimov Aziz" and row[2] == "+998909998877"  # the tutor's own number
    assert row[3] == ANSWERS["full_name"] and row[4] == PHONE
    assert row[6] == "DI-21" and row[7] == "3-kurs"
    assert row[8] == PASSPORT and row[9] == PINFL and row[10] == BIRTH
    assert row[11] == "O'zbekiston" and row[12] == "Toshkent shahri"
    assert row[17] == "Ishlamaydi" and row[22] == "Oila qurmagan"
    assert row[26] == "Nogironligi bor"
    assert row[27] == ANSWERS["father_name"] and row[30] == ANSWERS["mother_name"]
    assert row[33] == STUDENT

    # by group, and an empty group says so instead of sending a file
    await h.feed(callback_update(tutor_user, "tut:exf_pick:0:"))
    assert h.last_shown(TUTOR).text == texts.EXCEL_FULL_PICK_GROUP
    await h.feed(callback_update(tutor_user, f"tut:exf_group:{group.id}:"))
    assert len(h.documents(TUTOR)) == 2
    empty = await h.db.add_group(tutor.id, "DI-22")
    await h.feed(callback_update(tutor_user, f"tut:exf_group:{empty.id}:"))
    assert h.last_text(TUTOR) == texts.NO_DATA and len(h.documents(TUTOR)) == 2
    assert_all_callbacks_answered(h)


async def test_superadmin_exports_the_full_survey(h: Harness) -> None:
    tutor, group = await seed(h)
    await fill_full(h, make_user(STUDENT), tutor.id, group.id)
    admin = make_user(SUPERADMIN)
    h.clear()

    await h.feed(text_update(admin, "/admin"))
    assert inline_buttons(h.last_message(SUPERADMIN).reply_markup)["adm:excel_all_f:0"] == texts.BTN_ADMIN_EXCEL_FULL
    await h.feed(callback_update(admin, "adm:excel_all_f:0"))
    docs = h.documents(SUPERADMIN)
    assert len(docs) == 1
    wb = load_workbook(BytesIO(docs[0].document.data))
    assert wb.sheetnames == ["Barchasi", "Karimov Aziz"]
    assert [list(r) for r in wb["Barchasi"].iter_rows(values_only=True)][0] == list(FULL_HEADERS)

    await h.feed(callback_update(admin, "adm:list:0"), callback_update(admin, f"adm:view:{tutor.id}"))
    card = inline_buttons(h.last_shown(SUPERADMIN).reply_markup)
    assert card[f"adm:excel_tut_f:{tutor.id}"] == texts.BTN_EXCEL_FULL
    await h.feed(callback_update(admin, f"adm:excel_tut_f:{tutor.id}"))
    assert len(h.documents(SUPERADMIN)) == 2
    assert_all_callbacks_answered(h)


# ================================================================== tutor's phone


async def test_tutor_sets_their_own_phone(h: Harness) -> None:
    tutor, group = await seed(h)
    tutor_user = make_user(TUTOR)

    await h.feed(text_update(tutor_user, "/tutor"))
    assert inline_buttons(h.last_message(TUTOR).reply_markup)["tut:phone:0:"] == texts.BTN_TUTOR_PHONE
    await h.feed(callback_update(tutor_user, "tut:phone:0:"))
    assert h.last_text(TUTOR) == texts.TUTOR_ASK_PHONE

    await h.feed(text_update(tutor_user, "12345"))
    assert h.last_text(TUTOR) == texts.REG_PHONE_INVALID
    await h.feed(text_update(tutor_user, "+998909998877"))
    assert texts.TUTOR_PHONE_SAVED.format(phone="+998909998877") in h.texts_to(TUTOR)
    saved = await h.db.get_tutor(tutor.id)
    assert saved is not None and saved.phone == "+998909998877"
    assert await h.state_of(TUTOR) is None

    # opening it again shows the current number and accepts a shared contact
    await h.feed(text_update(tutor_user, "/phone"))
    assert "+998909998877" in h.last_text(TUTOR)
    await h.feed(contact_update(tutor_user, "998901112233", TUTOR))
    saved = await h.db.get_tutor(tutor.id)
    assert saved is not None and saved.phone == "+998901112233"
    assert_all_callbacks_answered(h)


# ================================================================== management


async def test_tutor_lists_and_deletes_full_survey_students(h: Harness) -> None:
    tutor, group = await seed(h)
    student = make_user(STUDENT)
    await fill_full(h, student, tutor.id, group.id)
    profile = await h.db.get_full_profile_by_telegram_id(STUDENT)
    assert profile is not None
    tutor_user = make_user(TUTOR)
    h.clear()

    await h.feed(callback_update(tutor_user, f"stu:svs:{group.id}:t:basic"))
    chooser = h.last_shown(TUTOR)
    assert "Qaysi anketa" in chooser.text
    assert inline_buttons(chooser.reply_markup) == {
        f"stu:list:{group.id}:t:basic": f"{texts.BTN_SURVEY_BASIC} — 0 ta",
        f"stu:list:{group.id}:t:full": f"{texts.BTN_SURVEY_FULL} — 1 ta",
        f"tut:view:{group.id}:": texts.BTN_BACK,
    }

    await h.feed(callback_update(tutor_user, f"stu:list:{group.id}:t:full"))
    listing = h.last_shown(TUTOR)
    assert texts.SURVEY_SHORT_LABELS["full"] in listing.text
    assert f"stu:view:{profile.id}:t:full" in inline_buttons(listing.reply_markup)

    await h.feed(callback_update(tutor_user, f"stu:view:{profile.id}:t:full"))
    card = h.last_shown(TUTOR)
    assert PASSPORT in card.text and PINFL in card.text  # the full card, not the basic one

    await h.feed(
        callback_update(tutor_user, f"stu:del:{profile.id}:t:full"),
        callback_update(tutor_user, f"stu:delok:{profile.id}:t:full"),
        callback_update(tutor_user, f"stu:bye_n:{profile.id}:t:full"),
    )
    assert await h.db.get_full_profile_by_telegram_id(STUDENT) is None
    assert h.messages(STUDENT) == []
    assert_all_callbacks_answered(h)


async def test_deleting_one_questionnaire_leaves_the_other_alone(h: Harness) -> None:
    from tests.test_flows import RegInput, register

    tutor, group = await seed(h)
    student = make_user(STUDENT, username="vali")
    await register(h, student, tutor.id, group.id, RegInput(full_name="Zokirov Vali"))
    await fill_full(h, student, tutor.id, group.id)
    basic = await h.db.get_student_by_telegram_id(STUDENT)
    assert basic is not None
    tutor_user = make_user(TUTOR)
    h.clear()

    await h.feed(
        callback_update(tutor_user, f"stu:del:{basic.id}:t:basic"),
        callback_update(tutor_user, f"stu:delok:{basic.id}:t:basic"),
        callback_update(tutor_user, f"stu:bye_n:{basic.id}:t:basic"),
    )
    assert await h.db.get_student_by_telegram_id(STUDENT) is None
    assert await h.db.get_full_profile_by_telegram_id(STUDENT) is not None  # untouched
    assert_all_callbacks_answered(h)
