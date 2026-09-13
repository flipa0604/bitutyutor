"""Excel builders: headers, row counts per sheet, residence labels, sheet-name sanitising."""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from bot.excel import (
    ALL_SHEET_NAME,
    HEADERS,
    build_all_tutors_workbook,
    build_group_workbook,
    build_residence_workbook,
    build_students_workbook,
    build_tutor_workbook,
    sanitize_sheet_name,
    unique_sheet_names,
)
from bot.models import Group, Student, Tutor


def make_student(idx: int, tutor: Tutor, group: Group, residence: str = "ttj", username: str | None = None) -> Student:
    return Student(
        id=idx,
        telegram_id=1000 + idx,
        username=username,
        tutor_id=tutor.id,
        group_id=group.id,
        full_name=f"Talaba {idx}",
        phone="+998901234567",
        direction="Dasturiy injiniring",
        residence=residence,
        address="TTJ" if residence == "ttj" else "Toshkent, Chilonzor 5",
        father_name="Ota Otayev",
        father_phone="+998901111111",
        mother_name="Ona Onayeva",
        mother_phone="+998902222222",
        created_at="2026-09-10 12:00:00",
        updated_at="2026-09-10 12:00:00",
        tutor_name=tutor.name,
        group_name=group.name,
    )


T1 = Tutor(id=1, name="Karimov Aziz", telegram_id=11, created_at="2026-01-01 00:00:00")
T2 = Tutor(id=2, name="Aliyev Bobur", telegram_id=22, created_at="2026-01-01 00:00:00")
G1 = Group(id=1, tutor_id=1, name="DI-21", created_at="2026-01-01 00:00:00")
G2 = Group(id=2, tutor_id=1, name="AI-22", created_at="2026-01-01 00:00:00")
G3 = Group(id=3, tutor_id=2, name="Barchasi", created_at="2026-01-01 00:00:00")


def load(data: bytes):
    assert data[:2] == b"PK"  # xlsx is a zip container
    return load_workbook(BytesIO(data))


def test_sanitize_sheet_name() -> None:
    assert sanitize_sheet_name("A[b]:c*d?e/f\\g") == "Abcdefg"
    assert sanitize_sheet_name("x" * 40) == "x" * 31
    assert sanitize_sheet_name("   ") == "Sheet"
    assert sanitize_sheet_name("'quoted'") == "quoted"


def _assert_valid_sheet_title(title: str) -> None:
    """Excel's rules: 1..31 chars, none of []:*?/\\, no leading/trailing apostrophe, not 'History'."""
    assert 1 <= len(title) <= 31, title
    assert not any(ch in title for ch in "[]:*?/\\"), title
    assert not title.startswith("'") and not title.endswith("'"), title
    assert title.lower() != "history", title


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        # 31st character is an apostrophe: truncating must not leave it at the end
        ("Abdullayev Abdulaziz Abdulla o'g'li", "Abdullayev Abdulaziz Abdulla o"),
        ("Yusupov Abdulla Abduraxmon o'g'li", "Yusupov Abdulla Abduraxmon o'g"),
        ("x" * 30 + "'y", "x" * 30),
        ("x" * 30 + " 'y", "x" * 30),
        ("x' '", "x"),
        ("''''''", "Sheet"),
        ("History", "History_"),
        ("history", "history_"),
    ],
)
def test_sanitize_sheet_name_after_truncation_and_reserved(name: str, expected: str) -> None:
    result = sanitize_sheet_name(name)
    assert result == expected
    _assert_valid_sheet_title(result)


def test_unique_sheet_names() -> None:
    names = unique_sheet_names(["Barchasi", "Barchasi", "barchasi", "a/b", "a\\b", "y" * 35, "y" * 35])
    assert names == ["Barchasi", "Barchasi (2)", "barchasi (3)", "ab", "ab (2)", "y" * 31, "y" * 27 + " (2)"]
    assert all(len(n) <= 31 for n in names)
    assert len({n.lower() for n in names}) == len(names)
    # the truncated base of a numbered duplicate is re-stripped too
    long_apostrophe = "a" * 26 + "'bcde"  # exactly 31 chars, apostrophe at index 26
    assert unique_sheet_names([long_apostrophe, long_apostrophe]) == [long_apostrophe, "a" * 26 + " (2)"]


def test_workbook_sheet_titles_for_long_uzbek_names_are_valid() -> None:
    tutor = Tutor(id=9, name="Abdullayev Abdulaziz Abdulla o'g'li", telegram_id=99, created_at="x")
    group = Group(id=9, tutor_id=9, name="Yusupov Abdulla Abduraxmon o'g'li guruhi", created_at="x")
    wb = load(build_all_tutors_workbook([tutor], [make_student(1, tutor, group)]))
    assert wb.sheetnames == [ALL_SHEET_NAME, "Abdullayev Abdulaziz Abdulla o"]
    wb = load(build_tutor_workbook([group], [make_student(1, tutor, group)]))
    assert wb.sheetnames == [ALL_SHEET_NAME, "Yusupov Abdulla Abduraxmon o'g"]
    for title in wb.sheetnames:
        _assert_valid_sheet_title(title)


def test_build_students_workbook_headers_rows_and_labels() -> None:
    students = [
        make_student(1, T1, G1, "ttj", username="first"),
        make_student(2, T1, G1, "kvartira"),
        make_student(3, T1, G2, "uy"),
    ]
    wb = load(build_students_workbook([("Sheet A", students), ("Empty", [])]))
    assert wb.sheetnames == ["Sheet A", "Empty"]

    ws = wb["Sheet A"]
    assert [c.value for c in ws[1]] == list(HEADERS)
    assert ws.max_row == 4
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref == "A1:P4"
    assert ws["A1"].font.bold

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert [r[0] for r in rows] == [1, 2, 3]
    assert [r[6] for r in rows] == ["TTJ", "Kvartira", "O'z uyi"]
    assert rows[0][1] == "Talaba 1"
    assert rows[0][4] == "Karimov Aziz" and rows[0][5] == "DI-21"
    assert rows[0][12] == "@first" and rows[1][12] is None
    assert rows[0][13] == 1001
    # phones stay text
    assert rows[0][2] == "+998901234567" and isinstance(rows[0][2], str)
    assert ws["C2"].number_format == "@"
    assert isinstance(rows[0][9], str) and isinstance(rows[0][11], str)

    empty = wb["Empty"]
    assert [c.value for c in empty[1]] == list(HEADERS)
    assert empty.max_row == 1


def test_build_students_workbook_without_sheets_has_headers_only() -> None:
    wb = load(build_students_workbook([]))
    assert wb.sheetnames == [ALL_SHEET_NAME]
    assert wb[ALL_SHEET_NAME].max_row == 1


def test_group_and_residence_workbooks() -> None:
    students = [make_student(1, T1, G1), make_student(2, T1, G1)]
    wb = load(build_group_workbook(G1, students))
    assert wb.sheetnames == ["DI-21"]
    assert wb["DI-21"].max_row == 3

    mixed = [make_student(3, T1, G2, "kvartira"), make_student(4, T1, G1, "kvartira")]
    wb = load(build_residence_workbook("kvartira", mixed))
    assert wb.sheetnames == ["Kvartira"]
    rows = list(wb["Kvartira"].iter_rows(min_row=2, values_only=True))
    assert [r[5] for r in rows] == ["AI-22", "DI-21"]  # Guruh column present, order preserved
    assert {r[6] for r in rows} == {"Kvartira"}


def test_tutor_workbook_has_all_sheet_plus_one_per_group() -> None:
    students = [make_student(1, T1, G1), make_student(2, T1, G1), make_student(3, T1, G2)]
    wb = load(build_tutor_workbook([G2, G1], students))
    assert wb.sheetnames == [ALL_SHEET_NAME, "AI-22", "DI-21"]
    assert wb[ALL_SHEET_NAME].max_row == 4
    assert wb["AI-22"].max_row == 2
    assert wb["DI-21"].max_row == 3


def test_all_tutors_workbook_unique_sheet_names() -> None:
    students = [make_student(1, T1, G1), make_student(2, T2, G3)]
    tutors = [
        T1,
        T2,
        Tutor(id=3, name="Barchasi", telegram_id=33, created_at="x"),
        Tutor(id=4, name="A/B", telegram_id=44, created_at="x"),
    ]
    wb = load(build_all_tutors_workbook(tutors, students))
    assert wb.sheetnames == [ALL_SHEET_NAME, "Karimov Aziz", "Aliyev Bobur", "Barchasi (2)", "AB"]
    assert wb[ALL_SHEET_NAME].max_row == 3
    assert wb["Karimov Aziz"].max_row == 2
    assert wb["Aliyev Bobur"].max_row == 2
    assert wb["Barchasi (2)"].max_row == 1
